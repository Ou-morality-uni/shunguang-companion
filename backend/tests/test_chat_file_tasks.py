import asyncio
import io
from threading import Event

import pytest
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessageChunk, ToolMessage

from backend.app import config, document_models, document_readers, document_tasks, main, runtime
from backend.app.filework import FileTaskRequest, FileWork, WorkError
from backend.app.routes import chat
from backend.app.store import Store
from backend.app.work_artifacts import artifact_bytes, attachment, attachment_bytes


@pytest.fixture
def setup(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'DATA_DIR', tmp_path / 'state')
    store = Store(config.DATA_DIR / 'test.sqlite3')
    store.initialize()
    source, output = tmp_path / 'source', tmp_path / 'output'
    source.mkdir()
    output.mkdir()
    work = FileWork(store)
    workspace = work.grant(str(source))
    store.set_output_workspace(work.grant(str(output))['id'])
    conversation = store.create_conversation()
    config.save_config(config.ModelConfig(llm_model='test', llm_api_key='dummy-key'))
    calls = []

    async def invoke(value, messages, check=lambda: None, tick=lambda: None):
        check()
        text = str(messages)
        calls.append(text)
        return (
            '本段保留终点标记。 [source｜来源位置]' if '这是分段整理' in text else '完整整理结果。 [source｜来源位置]'
        )

    monkeypatch.setattr(document_models, 'invoke', invoke)
    yield work, workspace['id'], source, output, conversation['id'], calls
    work.close()


def wait(work, task):
    return asyncio.run(work.wait(task['id']))


def test_file_tool_explains_invalid_fields_so_the_model_can_correct_and_retry(setup):
    from backend.app import feature_config
    from backend.app.tool_registry import ToolRegistry

    work, workspace_id, source, _, cid, _ = setup
    (source / 'source.txt').write_text('known source', encoding='utf-8')
    registry = ToolRegistry(feature_config.load(), work.store, {'conversation_id': cid}, [], work)
    args = {'workspace_id': workspace_id, 'steps': [{'action': 'read_document', 'path': 'source.txt', 'limit': 100}]}
    failed = asyncio.run(registry.execute('workspace_files', args))
    assert failed['ok'] is False
    assert failed['validation_errors'][0]['field'] == 'steps.0.limit'
    assert '50' in failed['message']
    args['steps'][0]['limit'] = 50
    corrected = asyncio.run(registry.execute('workspace_files', args))
    assert corrected['ok'] is True
    assert wait(work, corrected['task'])['status'] == 'complete'


def submit(setup, filename, scope='full', **kwargs):
    work, workspace_id, _, _, conversation_id, _ = setup
    return work.submit(
        FileTaskRequest(
            workspace_id=workspace_id,
            save_result=True,
            instruction='根据资料回答并标注来源',
            steps=[{'action': 'read_document', 'path': filename, 'scope': scope, **kwargs}],
        ),
        'chat',
        conversation_id,
    )


def test_new_summary_uses_project_outputs_without_changing_existing_files(setup):
    work, _, source, old_output, _, _ = setup
    (old_output / 'existing.md').write_text('keep this result', encoding='utf-8')
    (source / 'source.txt').write_text('source for automatic output', encoding='utf-8')
    task = wait(work, submit(setup, 'source.txt'))
    assert task['status'] == 'complete', task['error']
    artifact, data = artifact_bytes(work, task['id'])
    expected = config.DATA_DIR / 'outputs' / artifact['path']
    assert expected.is_file()
    assert expected.read_bytes() == data
    assert (old_output / 'existing.md').read_text(encoding='utf-8') == 'keep this result'


def test_ordinary_chat_reading_returns_a_summary_without_automatically_creating_a_file(setup):
    work, workspace_id, source, _, cid, _ = setup
    (source / 'source.txt').write_text('Known material', encoding='utf-8')
    task = wait(
        work,
        work.submit(
            FileTaskRequest(
                workspace_id=workspace_id,
                instruction='简要总结',
                steps=[{'action': 'read_document', 'path': 'source.txt', 'scope': 'full'}],
            ),
            'chat',
            cid,
        ),
    )
    assert task['status'] == 'complete', task['error']
    assert task['steps'][-1]['result']['abstract']
    assert not task['context'].get('artifact')
    assert not list(work.output_root.rglob('*.md'))


def test_chat_task_keeps_its_trigger_message_and_uploaded_sources_in_storage(setup):
    from backend.app import feature_config
    from backend.app.tool_registry import ToolRegistry

    work, _, _, _, cid, _ = setup
    uploaded = work.upload('source.txt', b'Known source')
    source = attachment(work, uploaded['task']['id'])
    message = work.store.add_message(cid, 'user', 'Read this document', metadata={'attachments': [source]})
    registry = ToolRegistry(feature_config.load(), work.store, {**message, 'conversation_id': cid}, [], work)
    result = asyncio.run(
        registry.execute(
            'workspace_files',
            {'workspace_id': source['workspace_id'], 'steps': [{'action': 'read_document', 'path': source['path']}]},
        )
    )
    completed = wait(work, result['task'])
    assert completed['context']['source_message_id'] == message['id']
    assert completed['context']['sources'] == [{'kind': 'attachment', 'id': source['id'], 'name': 'source.txt'}]
    stored = next(item for item in Store(work.store.path).list_file_tasks() if item['id'] == completed['id'])
    assert stored['context']['source_message_id'] == message['id']


def test_file_task_keeps_its_conversation_model_when_queued_and_global_defaults_change(setup, monkeypatch):
    work, _, source, _, cid, _ = setup
    config.save_config(
        config.ModelConfig(
            llm_base_url='https://selected.example/v1', llm_model='selected-model', llm_api_key='dummy-one'
        )
    )
    choice = {
        'profile_id': config.profile_id('https://selected.example/v1', 'openai'),
        'model': 'selected-model',
        'effort': 'default',
    }
    work.store.set_conversation_model(cid, choice)
    config.save_config(
        config.ModelConfig(
            llm_base_url='https://default.example/v1', llm_model='default-model', llm_api_key='dummy-default'
        )
    )
    started, release = Event(), Event()
    seen = []

    async def invoke(value, *args, **kwargs):
        seen.append((value.llm_base_url, value.llm_model))
        started.set()
        assert await asyncio.to_thread(release.wait, 5)
        return '资料整理完成。'

    monkeypatch.setattr(document_models, 'invoke', invoke)
    # Exercise multiple model calls without forty unrelated disk-reading batches.
    (source / 'source.txt').write_text('task material ' * 4000, encoding='utf-8')
    task = submit(setup, 'source.txt')
    try:
        assert started.wait(5)
        config.save_config(
            config.ModelConfig(
                llm_base_url='https://other.example/v1', llm_model='other-model', llm_api_key='dummy-two'
            )
        )
        work.store.set_conversation_model(
            cid,
            {
                'profile_id': config.profile_id('https://other.example/v1', 'openai'),
                'model': 'other-model',
                'effort': 'default',
            },
        )
    finally:
        release.set()
    result = wait(work, task)
    assert result['status'] == 'complete', result['error']
    assert len(seen) >= 2
    assert all(item == ('https://selected.example/v1', 'selected-model') for item in seen)


def test_summary_has_a_short_chat_abstract_and_saves_the_complete_markdown(setup, monkeypatch):
    import json

    work, _, source, _, _, _ = setup
    full = '# 完整整理\n\n' + '逐项核对的详细结果。' * 100

    async def invoke(*args, **kwargs):
        return json.dumps({'abstract': '资料有三项关键要求，详见完整成果。', 'summary': full}, ensure_ascii=False)

    monkeypatch.setattr(document_models, 'invoke', invoke)
    (source / 'source.txt').write_text('real task material', encoding='utf-8')
    task = wait(work, submit(setup, 'source.txt'))
    assert task['status'] == 'complete', task['error']
    result = task['steps'][-1]['result']
    assert result['summary'] == full
    assert result['abstract'] == '资料有三项关键要求，详见完整成果。'
    assert full in artifact_bytes(work, task['id'])[1].decode('utf-8')


def test_full_word_reads_later_blocks_and_summary_uses_all_chunks_then_saves_copy(setup):
    from docx import Document

    work, _, source, output, _, calls = setup
    document = Document()
    for index in range(65):
        document.add_paragraph(f'原文第{index}段 ' + '实际内容' * 250)
    document.add_paragraph('FINAL_UNIQUE_MARKER')
    document.save(source / 'source.docx')
    original = (source / 'source.docx').read_bytes()
    task = wait(work, submit(setup, 'source.docx'))
    assert task['status'] == 'complete', task['error']
    result = task['steps'][0]['result']
    assert len(result['blocks']) == 66 and not result['partial']
    assert any('FINAL_UNIQUE_MARKER' in call for call in calls)
    assert any('原文第0段' in call for call in calls)
    assert len(calls) > 1 and task['context']['budget']['calls'] == len(calls)
    artifact, data = artifact_bytes(work, task['id'])
    assert (config.DATA_DIR / 'outputs' / artifact['path']).read_bytes() == data
    assert '阅读范围与来源'.encode() in data and 'source.docx'.encode() in data
    assert (source / 'source.docx').read_bytes() == original


def test_full_excel_reads_beyond_old_row_and_column_caps(setup):
    from openpyxl import Workbook

    work, _, source, _, _, calls = setup
    book = Workbook()
    sheet = book.active
    sheet['A1'] = 'begin'
    sheet['DZ620'] = 'LAST_EXCEL_VALUE'
    book.save(source / 'source.xlsx')
    task = wait(work, submit(setup, 'source.xlsx'))
    assert task['status'] == 'complete', task['error']
    assert any('DZ620=LAST_EXCEL_VALUE' in block['text'] for block in task['steps'][0]['result']['blocks'])
    assert any('LAST_EXCEL_VALUE' in call for call in calls)
    assert not task['steps'][0]['result']['partial']


def test_full_single_long_block_continues_character_offsets_without_omissions(setup):
    work, _, source, _, _, _ = setup
    text = 'a' * 70000 + 'END_OF_LONG_LINE'
    (source / 'source.txt').write_text(text)
    task = wait(work, submit(setup, 'source.txt'))
    assert task['status'] == 'complete', task['error']
    assert ''.join(block['text'] for block in task['steps'][0]['result']['blocks']) == text
    assert not task['steps'][0]['result']['partial']


def test_workload_forecast_pauses_with_speech_and_resumes_the_same_cursor(setup, monkeypatch):
    work, _, source, _, _, calls = setup
    monkeypatch.setattr(document_tasks, 'ROUND_BATCHES', 1)
    (source / 'source.txt').write_text('\n'.join(f'line {index}' for index in range(100)))
    task = wait(work, submit(setup, 'source.txt'))
    assert task['status'] == 'waiting_limit' and calls == []
    assert len(task['steps'][0]['result']['blocks']) == 50
    assert task['events'][-1]['kind'] == 'waiting_limit' and task['events'][-1]['text']
    task = wait(work, work.continue_task(task['id'], True))
    assert task['status'] == 'complete'
    assert len(task['steps'][0]['result']['blocks']) == 100
    assert len(calls) == 1


def test_model_budget_continuation_reuses_saved_segment_notes_without_paid_retries(setup, monkeypatch):
    work, _, source, _, _, calls = setup
    monkeypatch.setattr(document_tasks, 'ROUND_CALLS', 1)
    (source / 'source.txt').write_text('真实长资料' * 18000, encoding='utf-8')
    task = wait(work, submit(setup, 'source.txt'))
    assert task['status'] == 'waiting_limit' and len(calls) == 1
    for _ in range(15):
        if task['status'] == 'complete':
            break
        task = wait(work, work.continue_task(task['id'], True))
    assert task['status'] == 'complete', task['error']
    assert len(calls) == len(set(calls))


def test_changed_source_cannot_be_silently_resumed_after_confirmation(setup, monkeypatch):
    work, _, source, _, _, calls = setup
    monkeypatch.setattr(document_tasks, 'ROUND_BATCHES', 1)
    path = source / 'source.txt'
    path.write_text('\n'.join(str(i) for i in range(100)))
    task = wait(work, submit(setup, 'source.txt'))
    path.write_text('changed')
    task = wait(work, work.continue_task(task['id'], True))
    assert task['status'] == 'failed' and calls == []
    assert len(task['steps'][0]['result']['blocks']) == 50


def test_focused_question_only_sends_matching_ranges_and_discloses_scope(setup):
    work, _, source, _, _, calls = setup
    lines = ['IRRELEVANT_PRIVATE_TEXT'] + [f'other {i}' for i in range(98)] + ['needle: target value']
    (source / 'source.txt').write_text('\n'.join(lines))
    task = wait(work, submit(setup, 'source.txt', scope='focused', query='needle'))
    assert task['status'] == 'complete'
    result = task['steps'][0]['result']
    assert result['partial'] and '关键词' in result['coverage']
    assert all('IRRELEVANT_PRIVATE_TEXT' not in call for call in calls)
    assert any('target value' in call for call in calls)


def test_unavailable_output_directory_preserves_summary_and_saving_does_not_call_model_again(setup):
    work, _, source, _, _, calls = setup
    blocked = config.DATA_DIR / 'outputs'
    blocked.write_text('directory temporarily unavailable', encoding='utf-8')
    (source / 'source.txt').write_text('one source')
    task = wait(work, submit(setup, 'source.txt'))
    assert task['status'] == 'waiting_output' and len(calls) == 1
    blocked.unlink()
    task = wait(work, work.continue_task(task['id'], True))
    assert task['status'] == 'complete' and len(calls) == 1
    assert artifact_bytes(work, task['id'])[1]


def test_downloads_revalidate_upload_and_output_fingerprints(setup):
    work, _, source, _, _, _ = setup
    uploaded = work.upload('source.txt', b'actual file')
    ref = attachment(work, uploaded['task']['id'])
    assert attachment_bytes(work, ref['id'])[1] == b'actual file'
    (work.upload_root / ref['path']).write_bytes(b'replaced')
    with pytest.raises(WorkError):
        attachment_bytes(work, ref['id'])
    (source / 'source.txt').write_text('source')
    task = wait(work, submit(setup, 'source.txt'))
    artifact, _ = artifact_bytes(work, task['id'])
    (work.root(artifact['workspace_id']) / artifact['path']).write_text('outside edit')
    with pytest.raises(WorkError):
        artifact_bytes(work, task['id'])


def test_queued_task_can_be_cancelled_before_the_worker_reaches_it(setup, monkeypatch):
    work, workspace_id, source, _, _, _ = setup
    entered, release = Event(), Event()
    original = work.operate

    def slow(root, step):
        entered.set()
        assert release.wait(5)
        return original(root, step)

    monkeypatch.setattr(work, 'operate', slow)
    first = work.submit(
        FileTaskRequest(workspace_id=workspace_id, steps=[{'action': 'write', 'path': 'one.txt', 'content': 'one'}])
    )
    assert entered.wait(3)
    second = work.submit(
        FileTaskRequest(workspace_id=workspace_id, steps=[{'action': 'write', 'path': 'two.txt', 'content': 'two'}])
    )
    assert work.cancel(second['id'])['status'] == 'cancelled'
    release.set()
    wait(work, first)
    assert wait(work, second)['status'] == 'cancelled'
    assert not (source / 'two.txt').exists()


def test_chat_attachments_are_persisted_and_background_task_does_not_block_next_turn(setup, monkeypatch):
    work, _, _, _, _, _ = setup
    monkeypatch.setattr(runtime, 'store', work.store)
    entered, release = Event(), Event()

    class Model:
        def bind_tools(self, definitions):
            return self

        async def astream(self, history):
            if isinstance(history[-1], ToolMessage):
                yield AIMessageChunk(content='任务已安排，咱们也可以继续聊。')
            elif '顺便聊聊' in str(history):
                yield AIMessageChunk(content='我在呢。')
            else:
                reference = next(
                    item
                    for message in work.store.messages(work.store.list_conversations()[0]['id'])
                    for item in message.get('attachments', [])
                )
                yield AIMessageChunk(
                    content='',
                    tool_calls=[
                        {
                            'name': 'workspace_files',
                            'id': 'read',
                            'args': {
                                'workspace_id': reference['workspace_id'],
                                'instruction': '总结整份',
                                'steps': [{'action': 'read_document', 'path': reference['path'], 'scope': 'full'}],
                            },
                        }
                    ],
                )

    monkeypatch.setattr(chat, 'build_model', lambda _: Model())
    original = document_tasks.read

    def slow(*args):
        entered.set()
        assert release.wait(5)
        return original(*args)

    monkeypatch.setattr(document_tasks, 'read', slow)
    with TestClient(main.app) as client:
        uploaded = client.post('/api/work/upload?name=source.txt', content=b'file content').json()
        response = client.post('/api/chat', json={'content': '总结整份资料', 'attachments': [uploaded['task']['id']]})
        assert response.status_code == 200 and 'event: done' in response.text
        assert entered.wait(3)
        conversation_id = work.store.list_conversations()[0]['id']
        assert work.store.messages(conversation_id)[0]['attachments'][0]['name'] == 'source.txt'
        assert (
            'event: done'
            in client.post('/api/chat', json={'conversation_id': conversation_id, 'content': '顺便聊聊'}).text
        )
        assert client.get(f'/api/work/tasks?conversation_id={conversation_id}&compact=true').json()
        release.set()
        task = work.store.list_file_tasks(conversation_id=conversation_id)[0]
        assert wait(runtime.work, task)['status'] == 'complete'


def test_attachment_preview_and_download_are_local_only_and_do_not_invoke_models(setup, monkeypatch):
    from PIL import Image

    work, _, _, _, _, calls = setup
    monkeypatch.setattr(runtime, 'store', work.store)
    output = io.BytesIO()
    Image.new('RGB', (40, 40), 'blue').save(output, format='PNG')
    with TestClient(main.app) as client:
        uploaded = client.post('/api/work/upload?name=example.png', content=output.getvalue()).json()
        task_id = uploaded['task']['id']
        assert client.get(f'/api/work/files/{task_id}').content == output.getvalue()
        assert client.get(f'/api/work/files/{task_id}/preview').status_code == 200
        preview = client.get(f'/api/work/files/{task_id}/image')
        assert preview.headers['content-type'] == 'image/jpeg' and calls == []
    with TestClient(main.app, client=('192.168.1.5', 123)) as client:
        assert client.get(f'/api/work/files/{task_id}').status_code == 403
        assert client.post('/api/chat', json={'content': 'read', 'attachments': [task_id]}).status_code == 403


def test_visual_overflow_can_be_paged_without_losing_text_or_images():
    from docx import Document
    from PIL import Image

    document = Document()
    document.add_paragraph('before')
    for index in range(10):
        image = io.BytesIO()
        Image.new('RGB', (4, 4), (index, 0, 0)).save(image, format='PNG')
        document.add_picture(io.BytesIO(image.getvalue()))
    output = io.BytesIO()
    document.save(output)
    first, images = document_readers.extract(output.getvalue(), 'source.docx', 1, 20, full=True)
    assert len(images) == 8 and first['next_image'] == 8
    second, remaining = document_readers.extract(output.getvalue(), 'source.docx', 1, 20, full=True, image_start=8)
    assert len(remaining) == 2 and 'next_image' not in second


def test_full_file_vision_consent_covers_later_image_batches_without_duplicate_calls(setup, monkeypatch):
    from docx import Document
    from PIL import Image

    work, _, source, _, _, _ = setup
    document = Document()
    for index in range(10):
        image = io.BytesIO()
        Image.new('RGB', (4, 4), (index, 0, 0)).save(image, format='PNG')
        document.add_picture(io.BytesIO(image.getvalue()))
    document.save(source / 'images.docx')
    task = wait(work, submit(setup, 'images.docx'))
    assert task['status'] == 'waiting_vision'
    assert '全部待识别图像' in task['steps'][0]['vision_request']['scope']
    config.save_config(config.ModelConfig(llm_model='vision', llm_api_key='dummy-key', llm_vision='supported'))
    profile = document_models.profiles()[0]
    requests = []

    async def recognize(value, image, check, tick):
        requests.append(image['location'])
        return '图像内容：' + image['location']

    monkeypatch.setattr(document_models, 'recognize', recognize)
    task = wait(work, work.approve_vision(task['id'], True, profile['id'], profile['signature']))
    assert task['status'] == 'complete', task['error']
    assert len(requests) == len(set(requests)) == 10
    assert len([block for block in task['steps'][0]['result']['blocks'] if block['kind'] == 'vision']) == 10
    assert not task['steps'][0]['result']['partial']
