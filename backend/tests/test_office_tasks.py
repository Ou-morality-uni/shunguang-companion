import asyncio
import io

from fastapi.testclient import TestClient

from backend.app import main, runtime
from backend.app.store import Store


def test_office_capabilities_explain_missing_engine_without_disabling_reading(tmp_path, monkeypatch):
    store = Store(tmp_path / 'office.sqlite3')
    store.initialize()
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app, raise_server_exceptions=True) as client:
        response = client.get('/api/work/capabilities')
    assert response.status_code == 200
    result = response.json()
    assert result['formats']['docx']['read'] is True
    assert result['formats']['xlsx']['read'] is True
    assert result['office']['install_url'] == 'https://www.libreoffice.org/download/download-libreoffice/'
    assert isinstance(result['office']['ready'], bool)
    if not result['office']['ready']:
        assert result['formats']['xlsx']['recalculate'] is False
        assert result['formats']['docx']['page_preview'] is False
        assert result['office']['message']


def test_project_local_office_is_detected_without_user_path_configuration(tmp_path, monkeypatch):
    import pytest

    from backend.app import config

    if not (config.ROOT / 'tools/LibreOffice/program/soffice.com').is_file():
        pytest.skip('Project-local LibreOffice is not installed')
    monkeypatch.setenv('PATH', '')
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        assert client.get('/api/work/capabilities').json()['office']['ready'] is True


def test_chat_can_create_and_download_a_real_word_report(tmp_path, monkeypatch):
    from docx import Document

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        conversation = store.create_conversation()
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': '调研报告.docx',
                'conversation_id': conversation['id'],
                'spec': {
                    'blocks': [
                        {'type': 'heading', 'text': '项目调研报告', 'level': 0},
                        {'type': 'paragraph', 'text': '根据已提供的材料整理。'},
                        {'type': 'table', 'header': ['项目', '数量'], 'rows': [['已确认需求', '3']]},
                    ]
                },
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        response = client.get('/api/work/artifacts/' + task['id'])
        document = Document(io.BytesIO(response.content))
        assert document.paragraphs[0].text == '项目调研报告'
        assert document.tables[0].cell(1, 1).text == '3'
        assert task['context']['artifact']['name'] == '调研报告.docx'
        preview = client.get('/api/work/artifacts/' + task['id'] + '/preview')
        assert preview.status_code == 200
        assert preview.json()['blocks'][0]['text'] == '项目调研报告'


def test_excel_task_keeps_numbers_and_delivers_a_native_chart(tmp_path, monkeypatch):
    import zipfile

    from openpyxl import load_workbook

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'xlsx',
                'name': '活动统计.xlsx',
                'spec': {
                    'sheets': [
                        {
                            'name': '活动统计',
                            'rows': [['项目', '人数'], ['跑步', 12], ['游泳', 8]],
                            'charts': [{'type': 'bar', 'data': 'B1:B3', 'categories': 'A2:A3', 'title': '活动人数'}],
                        }
                    ]
                },
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        content = client.get('/api/work/artifacts/' + task['id']).content
        workbook = load_workbook(io.BytesIO(content))
        assert workbook['活动统计']['B2'].value == 12
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            assert 'xl/charts/chart1.xml' in archive.namelist()


def test_excel_formulas_are_actually_calculated_and_recalculate_after_a_targeted_edit(tmp_path, monkeypatch):
    import pytest
    from openpyxl import load_workbook

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        if not client.get('/api/work/capabilities').json()['office']['ready']:
            pytest.skip('Real LibreOffice required for formula calculation')
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'xlsx',
                'name': 'analysis.xlsx',
                'spec': {
                    'sheets': [
                        {
                            'name': 'Data',
                            'rows': [['Item', 'Value'], ['A', 12], ['B', 8]],
                            'cells': {'B4': {'formula': '=SUM(B2:B3)'}},
                            'charts': [{'type': 'bar', 'data': 'B1:B3', 'categories': 'A2:A3', 'anchor': 'D1'}],
                        }
                    ]
                },
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        original = client.get('/api/work/artifacts/' + task['id']).content
        assert load_workbook(io.BytesIO(original), data_only=True).active['B4'].value == 20
        assert task['context']['artifact']['validation']['calculation'] == 'passed'
        resource = client.get('/api/work/resources/artifact/' + task['id'] + '/targets').json()
        revised = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'xlsx',
                'name': 'revised.xlsx',
                'operation': 'edit',
                'source': {'kind': 'artifact', 'id': task['id'], 'version': resource['version']},
                'spec': {'edits': [{'target': 's:0:B2', 'old': 12, 'value': 15}]},
            },
        )
        result = asyncio.run(runtime.work.wait(revised.json()['id']))
        assert result['status'] == 'complete', result['error']
        data = client.get('/api/work/artifacts/' + result['id']).content
        assert load_workbook(io.BytesIO(data), data_only=True).active['B4'].value == 23
        book = load_workbook(io.BytesIO(data), data_only=False)
        assert book.active['B4'].value == '=SUM(B2:B3)'
        import zipfile

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            assert any(name.startswith('xl/charts/chart') and name.endswith('.xml') for name in archive.namelist())
        assert client.get('/api/work/artifacts/' + task['id']).content == original


def test_excel_external_conditional_formulas_are_rejected_before_publication(tmp_path, monkeypatch):
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'xlsx',
                'name': 'unsafe.xlsx',
                'spec': {
                    'sheets': [
                        {
                            'name': 'Data',
                            'rows': [[1]],
                            'conditional_formats': [
                                {
                                    'type': 'cell_is',
                                    'range': 'A1:A1',
                                    'formula': ['WEBSERVICE ("https://blocked.invalid/data")'],
                                }
                            ],
                        }
                    ]
                },
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'failed'
        assert '公式' in task['error']
        assert not task['context'].get('artifacts')


def test_chat_routes_office_tools_without_granting_access_to_another_conversation(tmp_path, monkeypatch):
    from backend.app import feature_config
    from backend.app.tool_registry import ToolRegistry

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app):
        first = store.create_conversation()
        second = store.create_conversation()
        registry = ToolRegistry(feature_config.load(), store, {'conversation_id': first['id']}, [], runtime.work)
        guide = asyncio.run(registry.execute('office_guide', {'format': 'docx'}))
        assert guide['ok'] is True
        assert 'blocks' in guide['guide']
        result = asyncio.run(
            registry.execute(
                'office_files',
                {
                    'format': 'docx',
                    'name': '报告.docx',
                    'conversation_id': second['id'],
                    'spec': {'blocks': [{'type': 'paragraph', 'text': '只属于当前聊天'}]},
                },
            )
        )
        assert result['ok'] is False
        result = asyncio.run(
            registry.execute(
                'office_files',
                {
                    'format': 'docx',
                    'name': '报告.docx',
                    'spec': {'blocks': [{'type': 'paragraph', 'text': '只属于当前聊天'}]},
                },
            )
        )
        assert result['ok'] is True
        task = asyncio.run(runtime.work.wait(result['task']['id']))
        assert task['conversation_id'] == first['id']


def test_pdf_task_preserves_chinese_text_and_does_not_interpret_markup(tmp_path, monkeypatch):
    import pypdfium2 as pdfium

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'pdf',
                'name': '中文报告.pdf',
                'spec': {
                    'elements': [
                        {'type': 'heading', 'text': '活动统计报告'},
                        {'type': 'paragraph', 'text': '人数合计二十人。<b>原文标记</b>'},
                    ]
                },
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        content = client.get('/api/work/artifacts/' + task['id']).content
        with pdfium.PdfDocument(content) as pdf:
            page = pdf[0]
            textpage = page.get_textpage()
            try:
                text = textpage.get_text_range()
                assert '活动统计报告' in text
                assert '<b>原文标记</b>' in text
            finally:
                textpage.close()
                page.close()
        preview = client.get('/api/work/resources/artifact/' + task['id'] + '/pages')
        assert preview.status_code == 200
        assert preview.json()['pages'] == 1
        image = client.get('/api/work/resources/artifact/' + task['id'] + '/image?page=1')
        assert image.status_code == 200
        assert image.content.startswith(b'\xff\xd8')


def test_uploaded_word_exposes_versioned_targets_for_duplicate_text(tmp_path, monkeypatch):
    from docx import Document

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    source = Document()
    source.add_paragraph('同名内容')
    source.add_paragraph('同名内容')
    data = io.BytesIO()
    source.save(data)
    with TestClient(main.app) as client:
        upload = client.post('/api/work/upload?name=source.docx', content=data.getvalue())
        assert upload.status_code == 200
        resource_id = upload.json()['task']['id']
        result = client.get(f'/api/work/resources/attachment/{resource_id}/targets')
        assert result.status_code == 200, result.text
        value = result.json()
        assert len(value['version']) == 64
        targets = value['targets']
        assert targets[0]['text'] == targets[1]['text'] == '同名内容'
        assert targets[0]['id'] != targets[1]['id']


def test_large_document_target_preview_is_bounded_and_discloses_truncation(tmp_path, monkeypatch):
    from docx import Document

    document = Document()
    for _ in range(8):
        document.add_paragraph('x' * 20000)
    source = io.BytesIO()
    document.save(source)
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=large.docx', content=source.getvalue()).json()['task']['id']
        response = client.get(f'/api/work/resources/attachment/{identity}/targets')
        assert response.status_code == 200
        result = response.json()
        assert result['truncated'] is True
        assert sum(len(target['text']) for target in result['targets']) <= 120000


def test_invalid_office_upload_has_a_readable_preview_error(tmp_path, monkeypatch):
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app, raise_server_exceptions=False) as client:
        identity = client.post('/api/work/upload?name=broken.docx', content=b'not an office archive').json()['task'][
            'id'
        ]
        response = client.get(f'/api/work/resources/attachment/{identity}/targets')
        assert response.status_code == 422, response.text
        assert '文件结构' in response.json()['detail']


def test_word_revision_changes_only_selected_paragraph_and_keeps_original(tmp_path, monkeypatch):
    from docx import Document

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    source = Document()
    source.add_paragraph('同名内容').runs[0].bold = True
    paragraph = source.add_paragraph()
    paragraph.add_run('同名').italic = True
    paragraph.add_run('内容')
    data = io.BytesIO()
    source.save(data)
    with TestClient(main.app) as client:
        original = data.getvalue()
        upload_id = client.post('/api/work/upload?name=source.docx', content=original).json()['task']['id']
        metadata = client.get(f'/api/work/resources/attachment/{upload_id}/targets').json()
        response = client.post(
            '/api/work/office/tasks',
            json={
                'operation': 'edit',
                'format': 'docx',
                'name': '修订.docx',
                'source': {'kind': 'attachment', 'id': upload_id, 'version': metadata['version']},
                'spec': {'edits': [{'target': 'p:1', 'old': '同名内容', 'text': '仅修改第二段'}]},
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        revised = Document(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert revised.paragraphs[0].text == '同名内容'
        assert revised.paragraphs[0].runs[0].bold is True
        assert revised.paragraphs[1].text == '仅修改第二段'
        assert revised.paragraphs[1].runs[0].italic is True
        assert client.get('/api/work/files/' + upload_id).content == original


def test_word_template_fills_blank_text_and_header_without_losing_layout(tmp_path, monkeypatch):
    from docx import Document
    from docx.shared import Mm

    template = Document()
    template.sections[0].left_margin = Mm(31)
    template.sections[0].header.paragraphs[0].text = 'Company template'
    template.add_paragraph('')
    template.add_paragraph('Keep this paragraph', style='Heading 1')
    buffer = io.BytesIO()
    template.save(buffer)
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=template.docx', content=buffer.getvalue()).json()['task']['id']
        resource = client.get(f'/api/work/resources/attachment/{identity}/targets').json()
        assert any(item['id'] == 'h:0:0' for item in resource['targets'])
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'filled.docx',
                'source': {'kind': 'attachment', 'id': identity, 'version': resource['version']},
                'spec': {
                    'edits': [
                        {'target': 'p:0', 'old': '', 'text': 'New body'},
                        {'target': 'h:0:0', 'old': 'Company', 'text': 'Project'},
                    ]
                },
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        document = Document(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert document.paragraphs[0].text == 'New body'
        assert document.paragraphs[1].text == 'Keep this paragraph'
        assert document.paragraphs[1].style.name == 'Heading 1'
        assert document.sections[0].header.paragraphs[0].text == 'Project template'
        assert document.sections[0].left_margin.twips == 1757  # 31 mm, rounded to Word's 1/20 pt unit.


def test_word_embeds_an_explicit_image_resource_as_an_editable_document_object(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image

    image = io.BytesIO()
    Image.new('RGBA', (200, 100), (230, 30, 60, 120)).save(image, format='PNG')
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=figure.png', content=image.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'illustrated.docx',
                'spec': {
                    'blocks': [
                        {'type': 'paragraph', 'text': 'Native text next to the figure.'},
                        {
                            'type': 'image',
                            'image': {'kind': 'attachment', 'id': identity, 'version': version},
                            'width_cm': 8,
                        },
                    ]
                },
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        document = Document(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert document.paragraphs[0].text == 'Native text next to the figure.'
        assert len(document.inline_shapes) == 1
        assert document.inline_shapes[0].width == 2 * document.inline_shapes[0].height


def test_presentation_and_pdf_embed_only_the_requested_image_resource(tmp_path, monkeypatch):
    import zipfile

    import pypdfium2 as pdfium
    from PIL import Image
    from pptx import Presentation

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    image = io.BytesIO()
    Image.new('RGB', (200, 100), 'red').save(image, format='PNG')
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=figure.png', content=image.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']
        block = {'type': 'image', 'image': {'kind': 'attachment', 'id': identity, 'version': version}}
        for kind, spec in [
            ('pptx', {'slides': [{'title': 'Image and editable text', 'blocks': [block]}]}),
            ('pdf', {'elements': [{'type': 'heading', 'text': 'Illustrated PDF'}, block]}),
        ]:
            response = client.post(
                '/api/work/office/tasks', json={'format': kind, 'name': 'illustrated.' + kind, 'spec': spec}
            )
            task = asyncio.run(runtime.work.wait(response.json()['id']))
            assert task['status'] == 'complete', task['error']
            output = client.get('/api/work/artifacts/' + task['id']).content
            if kind == 'pptx':
                deck = Presentation(io.BytesIO(output))
                assert any(
                    shape.has_text_frame and shape.text == 'Image and editable text' for shape in deck.slides[0].shapes
                )
                with zipfile.ZipFile(io.BytesIO(output)) as archive:
                    assert any(name.startswith('ppt/media/') for name in archive.namelist())
            else:
                with pdfium.PdfDocument(output) as document:
                    page = document[0]
                    assert any(page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_IMAGE]))
                    page.close()


def test_background_generator_cannot_embed_an_image_not_attached_to_its_chat(tmp_path, monkeypatch):
    import json

    from PIL import Image

    from backend.app import config, document_models, feature_config
    from backend.app.tool_registry import ToolRegistry

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    config.save_config(config.ModelConfig(llm_model='test', llm_api_key='dummy'))
    image = io.BytesIO()
    Image.new('RGB', (20, 10), 'red').save(image, format='PNG')
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=other.png', content=image.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']

        async def generate(*args, **kwargs):
            return json.dumps(
                {'blocks': [{'type': 'image', 'image': {'kind': 'attachment', 'id': identity, 'version': version}}]}
            )

        monkeypatch.setattr(document_models, 'invoke', generate)
        conversation = store.create_conversation()
        registry = ToolRegistry(feature_config.load(), store, {'conversation_id': conversation['id']}, [], runtime.work)
        result = asyncio.run(
            registry.execute(
                'office_files', {'format': 'docx', 'name': 'private.docx', 'instruction': 'Make a report.'}
            )
        )
        assert result['ok'] is True
        task = asyncio.run(runtime.work.wait(result['task']['id']))
        assert task['status'] == 'failed'
        assert '聊天' in task['error']
        assert not task['context'].get('artifacts')


def test_pdf_revision_delivers_a_word_copy_and_keeps_the_pdf_source(tmp_path, monkeypatch):
    from docx import Document
    from reportlab.pdfgen.canvas import Canvas

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    original = io.BytesIO()
    canvas = Canvas(original)
    canvas.drawString(50, 700, 'Original PDF report')
    canvas.save()
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=source.pdf', content=original.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'revision.docx',
                'operation': 'edit',
                'source': {'kind': 'attachment', 'id': identity, 'version': version},
                'spec': {
                    'blocks': [
                        {'type': 'heading', 'text': 'Revised report', 'level': 1},
                        {'type': 'paragraph', 'text': 'Corrected paragraph. [source.pdf | page 1]'},
                    ]
                },
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        result = Document(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert result.paragraphs[1].text == 'Corrected paragraph. [source.pdf | page 1]'
        assert '排版' in task['steps'][-1]['result']['abstract']
        assert client.get('/api/work/files/' + identity).content == original.getvalue()


def test_word_task_delivers_editable_source_and_pdf_as_separate_downloads(tmp_path, monkeypatch):
    import pytest

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        if not client.get('/api/work/capabilities').json()['office']['ready']:
            pytest.skip('Real LibreOffice required')
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'report.docx',
                'export_pdf': True,
                'spec': {'blocks': [{'type': 'paragraph', 'text': 'The complete report.'}]},
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        assert [item['name'] for item in task['context']['artifacts']] == ['report.docx', 'report.pdf']
        assert client.get('/api/work/artifacts/' + task['id']).content.startswith(b'PK')
        assert client.get('/api/work/artifacts/' + task['id'] + '~1').content.startswith(b'%PDF-')


def test_complex_office_task_waits_for_outline_confirmation_before_creating_files(tmp_path, monkeypatch):
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'complex.docx',
                'outline': '背景、数据分析、结论三部分；确认后制作。',
                'spec': {'blocks': [{'type': 'paragraph', 'text': 'Approved report.'}]},
            },
        )
        assert response.status_code == 201, response.text
        task = response.json()
        assert task['status'] == 'waiting_limit'
        assert task['context']['pause']['kind'] == 'outline'
        assert client.get('/api/work/artifacts/' + task['id']).status_code == 404
        assert client.post('/api/work/tasks/' + task['id'] + '/continue', json={'approved': True}).status_code == 200
        completed = asyncio.run(runtime.work.wait(task['id']))
        assert completed['status'] == 'complete', completed['error']
        assert client.get('/api/work/artifacts/' + task['id']).content.startswith(b'PK')
        assert client.post('/api/work/tasks/' + task['id'] + '/continue', json={'approved': True}).status_code == 409


def test_failed_office_save_can_resume_without_another_model_call(tmp_path, monkeypatch):
    import json
    import tempfile

    from backend.app import config, document_models

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    config.save_config(config.ModelConfig(llm_model='test', llm_api_key='dummy'))
    calls = []

    async def generate(*args, **kwargs):
        calls.append(True)
        return json.dumps({'blocks': [{'type': 'paragraph', 'text': 'Saved after retry.'}]})

    monkeypatch.setattr(document_models, 'invoke', generate)
    mkstemp = tempfile.mkstemp
    blocked = [True]

    def unavailable(*args, **kwargs):
        if blocked[0] and kwargs.get('prefix') == '.shunguang-':
            raise PermissionError('Synthetic file lock')
        return mkstemp(*args, **kwargs)

    monkeypatch.setattr(tempfile, 'mkstemp', unavailable)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'retry.docx',
                'instruction': 'Make a short report.',
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'waiting_output', task['error']
        assert not task['context'].get('artifacts')
        blocked[0] = False
        assert client.post('/api/work/tasks/' + task['id'] + '/continue', json={'approved': True}).status_code == 200
        task = asyncio.run(runtime.work.wait(task['id']))
        assert task['status'] == 'complete', task['error']
        assert len(calls) == 1
        assert client.get('/api/work/artifacts/' + task['id']).content.startswith(b'PK')


def test_explicit_overwrite_waits_for_confirmation_and_can_restore_the_original(tmp_path, monkeypatch):
    import hashlib

    from docx import Document

    from backend.app import config

    monkeypatch.setattr(config, 'DATA_DIR', tmp_path / 'state')
    original = Document()
    original.add_paragraph('Original text')
    folder = tmp_path / 'authorized'
    folder.mkdir()
    destination = folder / 'original.docx'
    original.save(destination)
    original_bytes = destination.read_bytes()
    store = Store(config.DATA_DIR / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        grant = client.post('/api/work/workspaces', json={'path': str(folder)})
        assert grant.status_code == 201, grant.text
        workspace = grant.json()
        identity = client.post('/api/work/upload?name=original.docx', content=original_bytes).json()['task']['id']
        version = hashlib.sha256(original_bytes).hexdigest()
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'revision.docx',
                'operation': 'edit',
                'source': {'kind': 'attachment', 'id': identity, 'version': version},
                'overwrite': {'workspace_id': workspace['id'], 'path': 'original.docx', 'version': version},
                'spec': {'edits': [{'target': 'p:0', 'old': 'Original', 'text': 'Revised'}]},
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'waiting_limit'
        assert task['context']['pause']['kind'] == 'overwrite'
        assert destination.read_bytes() == original_bytes
        assert client.post('/api/work/tasks/' + task['id'] + '/continue', json={'approved': True}).status_code == 200
        task = asyncio.run(runtime.work.wait(task['id']))
        assert task['status'] == 'complete', task['error']
        assert task['can_restore'] is True
        assert Document(destination).paragraphs[0].text == 'Revised text'
        assert client.post('/api/work/tasks/' + task['id'] + '/restore').status_code == 200
        restored = asyncio.run(runtime.work.wait(task['id']))
        assert restored['status'] == 'restored', restored['error']
        assert destination.read_bytes() == original_bytes


def test_chat_can_import_an_authorized_template_with_a_versioned_overwrite_target(tmp_path, monkeypatch):
    import hashlib

    from docx import Document

    from backend.app import config, feature_config
    from backend.app.tool_registry import ToolRegistry

    monkeypatch.setattr(config, 'DATA_DIR', tmp_path / 'state')
    store = Store(config.DATA_DIR / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    folder = tmp_path / 'sources'
    folder.mkdir()
    document = Document()
    document.add_paragraph('Authorized template')
    document.save(folder / 'template.docx')
    with TestClient(main.app) as client:
        workspace = client.post('/api/work/workspaces', json={'path': str(folder)}).json()
        cid = store.create_conversation()['id']
        registry = ToolRegistry(feature_config.load(), store, {'conversation_id': cid}, [], runtime.work)
        result = asyncio.run(
            registry.execute('office_import', {'workspace_id': workspace['id'], 'path': 'template.docx'})
        )
        assert result['ok'] is True, result
        assert result['original']['version'] == hashlib.sha256((folder / 'template.docx').read_bytes()).hexdigest()
        ref = result['source']
        inspected = asyncio.run(registry.execute('office_inspect', {'kind': ref['kind'], 'id': ref['id']}))
        assert inspected['ok'] is True
        assert inspected['resource']['targets'][0]['text'] == 'Authorized template'
        forbidden = asyncio.run(
            registry.execute('office_import', {'workspace_id': workspace['id'], 'path': '../state/office.sqlite3'})
        )
        assert forbidden['ok'] is False


def test_office_task_can_resume_authorized_vision_for_an_uploaded_source(tmp_path, monkeypatch):
    import json

    from PIL import Image

    from backend.app import config, document_models

    config.save_config(config.ModelConfig(llm_model='text', llm_api_key='dummy', llm_vision='unsupported'))
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    image = io.BytesIO()
    Image.new('RGB', (10, 10), 'blue').save(image, format='PNG')

    async def recognize(*args, **kwargs):
        return 'Recognized blue image'

    async def generate(_value, messages, check=lambda: None, tick=lambda: None):
        assert 'Recognized blue image' in str(messages)
        return json.dumps({'blocks': [{'type': 'paragraph', 'text': 'Recognized source.'}]})

    monkeypatch.setattr(document_models, 'recognize', recognize)
    monkeypatch.setattr(document_models, 'invoke', generate)
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=source.png', content=image.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': 'image-report.docx',
                'instruction': 'Describe this source.',
                'sources': [{'kind': 'attachment', 'id': identity, 'version': version}],
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'waiting_vision'
        config.save_config(config.ModelConfig(llm_model='vision', llm_api_key='dummy', llm_vision='supported'))
        profile = document_models.profiles()[0]
        response = client.post(
            '/api/work/tasks/' + task['id'] + '/vision',
            json={'approved': True, 'profile_id': profile['id'], 'signature': profile['signature']},
        )
        assert response.status_code == 200, response.text
        task = asyncio.run(runtime.work.wait(task['id']))
        assert task['status'] == 'complete', task['error']


def test_presentation_delivers_editable_text_table_and_chart(tmp_path, monkeypatch):
    from pptx import Presentation

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'pptx',
                'name': '活动汇报.pptx',
                'spec': {
                    'slides': [
                        {'title': '活动汇报', 'blocks': [{'type': 'paragraph', 'text': '本周有二十人参加。'}]},
                        {
                            'title': '活动数据',
                            'blocks': [
                                {'type': 'table', 'rows': [['项目', '人数'], ['跑步', '12'], ['游泳', '8']]},
                                {
                                    'type': 'chart',
                                    'categories': ['跑步', '游泳'],
                                    'series': [{'name': '人数', 'values': [12, 8]}],
                                },
                            ],
                        },
                    ]
                },
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        deck = Presentation(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert len(deck.slides) == 2
        assert any(shape.has_text_frame and shape.text == '活动汇报' for shape in deck.slides[0].shapes)
        assert any(shape.has_table for shape in deck.slides[1].shapes)
        assert any(shape.has_chart for shape in deck.slides[1].shapes)


def test_office_export_reports_missing_engine_or_delivers_real_pdf(tmp_path, monkeypatch):
    from docx import Document

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    document = Document()
    document.add_paragraph('PDF 导出测试')
    data = io.BytesIO()
    document.save(data)
    with TestClient(main.app) as client:
        available = client.get('/api/work/capabilities').json()['office']['ready']
        resource_id = client.post('/api/work/upload?name=source.docx', content=data.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{resource_id}/targets').json()['version']
        response = client.post(
            '/api/work/office/tasks',
            json={
                'operation': 'export',
                'format': 'pdf',
                'name': '导出.pdf',
                'spec': {},
                'source': {'kind': 'attachment', 'id': resource_id, 'version': version},
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        if available:
            assert task['status'] == 'complete', task['error']
            assert client.get('/api/work/artifacts/' + task['id']).content.startswith(b'%PDF-')
        else:
            assert task['status'] == 'failed'
            assert 'LibreOffice' in task['error']
            assert not task['context'].get('artifact')


def test_standard_ppt_can_preview_and_export_with_its_print_settings_preserved(tmp_path, monkeypatch):
    import pytest
    from pptx import Presentation

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = 'Native presentation'
    original = io.BytesIO()
    deck.save(original)
    with TestClient(main.app) as client:
        if not client.get('/api/work/capabilities').json()['office']['ready']:
            pytest.skip('Real LibreOffice required for presentation rendering')
        identity = client.post('/api/work/upload?name=slides.pptx', content=original.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']
        pages = client.get(f'/api/work/resources/attachment/{identity}/pages?version={version}')
        assert pages.status_code == 200, pages.text
        assert pages.json()['pages'] == 1
        image = client.get(f'/api/work/resources/attachment/{identity}/image?version={version}')
        assert image.status_code == 200 and image.content.startswith(b'\xff\xd8')
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'pdf',
                'name': 'slides.pdf',
                'operation': 'export',
                'source': {'kind': 'attachment', 'id': identity, 'version': version},
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        assert client.get('/api/work/artifacts/' + task['id']).content.startswith(b'%PDF-')
        assert client.get('/api/work/files/' + identity).content == original.getvalue()


def test_image_task_crops_and_resizes_a_copy_without_losing_alpha(tmp_path, monkeypatch):
    from PIL import Image

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    image = Image.new('RGBA', (20, 10), (100, 40, 20, 90))
    data = io.BytesIO()
    image.save(data, format='PNG')
    with TestClient(main.app) as client:
        original = data.getvalue()
        resource_id = client.post('/api/work/upload?name=source.png', content=original).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{resource_id}/targets').json()['version']
        response = client.post(
            '/api/work/office/tasks',
            json={
                'operation': 'edit',
                'format': 'png',
                'name': '裁剪.png',
                'source': {'kind': 'attachment', 'id': resource_id, 'version': version},
                'spec': {'crop': [0, 0, 10, 10], 'size': [5, 5]},
            },
        )
        assert response.status_code == 201, response.text
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        with Image.open(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content)) as result:
            assert result.size == (5, 5)
            assert result.getpixel((2, 2))[3] == 90
        assert client.get('/api/work/files/' + resource_id).content == original


def test_chat_inspection_uses_attached_resource_and_rejects_another_chat(tmp_path, monkeypatch):
    from docx import Document

    from backend.app import feature_config
    from backend.app.tool_registry import ToolRegistry

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    document = Document()
    document.add_paragraph('可用于精确修改的内容')
    data = io.BytesIO()
    document.save(data)
    with TestClient(main.app) as client:
        upload = client.post('/api/work/upload?name=source.docx', content=data.getvalue()).json()
        first, second = store.create_conversation(), store.create_conversation()
        store.begin_turn(first['id'], '看看这份文件', {'attachments': [{'id': upload['task']['id']}]})
        own = ToolRegistry(feature_config.load(), store, {'conversation_id': first['id']}, [], runtime.work)
        result = asyncio.run(own.execute('office_inspect', {'kind': 'attachment', 'id': upload['task']['id']}))
        assert result['ok'] is True
        assert result['resource']['targets'][0]['text'] == '可用于精确修改的内容'
        assert len(result['resource']['version']) == 64
        foreign = ToolRegistry(feature_config.load(), store, {'conversation_id': second['id']}, [], runtime.work)
        assert (
            asyncio.run(foreign.execute('office_inspect', {'kind': 'attachment', 'id': upload['task']['id']}))['ok']
            is False
        )


def test_ppt_revision_preserves_untouched_slide_bytes_and_editable_chart(tmp_path, monkeypatch):
    import zipfile

    from pptx import Presentation

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        created = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'pptx',
                'name': '原稿.pptx',
                'spec': {
                    'slides': [
                        {'title': '保持不变', 'blocks': []},
                        {
                            'title': '旧标题',
                            'blocks': [
                                {
                                    'type': 'chart',
                                    'categories': ['甲', '乙'],
                                    'series': [{'name': '人数', 'values': [12, 8]}],
                                }
                            ],
                        },
                    ]
                },
            },
        ).json()
        first = asyncio.run(runtime.work.wait(created['id']))
        original = client.get('/api/work/artifacts/' + first['id']).content
        resource = client.get('/api/work/resources/artifact/' + first['id'] + '/targets').json()
        target = next(item for item in resource['targets'] if item['text'] == '旧标题')
        response = client.post(
            '/api/work/office/tasks',
            json={
                'operation': 'edit',
                'format': 'pptx',
                'name': '修订.pptx',
                'source': {'kind': 'artifact', 'id': first['id'], 'version': resource['version']},
                'spec': {'edits': [{'target': target['id'], 'old': '旧标题', 'text': '新的标题'}]},
            },
        )
        assert response.status_code == 201
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        revised = client.get('/api/work/artifacts/' + task['id']).content
        deck = Presentation(io.BytesIO(revised))
        assert any(shape.has_text_frame and shape.text == '新的标题' for shape in deck.slides[1].shapes)
        assert any(shape.has_chart for shape in deck.slides[1].shapes)
        with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(io.BytesIO(revised)) as after:
            assert before.read('ppt/slides/slide1.xml') == after.read('ppt/slides/slide1.xml')


def test_ppt_chart_revision_updates_native_data_without_rewriting_untouched_slides(tmp_path, monkeypatch):
    import zipfile

    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = 'Keep this slide'
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    chart_data = CategoryChartData()
    chart_data.categories = ['A', 'B']
    chart_data.add_series('Count', [12, 8])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), chart_data)
    stream = io.BytesIO()
    deck.save(stream)
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=chart.pptx', content=stream.getvalue()).json()['task']['id']
        resource = client.get(f'/api/work/resources/attachment/{identity}/targets').json()
        target = next(item for item in resource['targets'] if item.get('object_type') == 'chart')
        assert target['chart']['series'][0]['values'] == [12, 8]
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'pptx',
                'name': 'updated.pptx',
                'operation': 'edit',
                'source': {'kind': 'attachment', 'id': identity, 'version': resource['version']},
                'spec': {
                    'edits': [
                        {
                            'target': target['id'],
                            'chart': {'categories': ['A', 'B'], 'series': [{'name': 'Count', 'values': [15, 8]}]},
                        }
                    ]
                },
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        output = client.get('/api/work/artifacts/' + task['id']).content
        changed = Presentation(io.BytesIO(output))
        assert list(changed.slides[1].shapes[0].chart.series[0].values) == [15, 8]
        with zipfile.ZipFile(io.BytesIO(output)) as revised, zipfile.ZipFile(io.BytesIO(stream.getvalue())) as original:
            assert revised.read('ppt/slides/slide1.xml') == original.read('ppt/slides/slide1.xml')
            assert any(name.startswith('ppt/embeddings/') and name.endswith('.xlsx') for name in revised.namelist())


def test_ppt_table_revision_changes_only_the_selected_cell(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    table = (
        deck.slides.add_slide(deck.slide_layouts[6])
        .shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(3))
        .table
    )
    table.cell(0, 0).text, table.cell(0, 1).text = 'Item', 'Count'
    table.cell(1, 0).text, table.cell(1, 1).text = 'A', '12'
    stream = io.BytesIO()
    deck.save(stream)
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=table.pptx', content=stream.getvalue()).json()['task']['id']
        resource = client.get(f'/api/work/resources/attachment/{identity}/targets').json()
        target = next(item for item in resource['targets'] if item['id'].endswith(':cell:1:1'))
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'pptx',
                'name': 'updated.pptx',
                'operation': 'edit',
                'source': {'kind': 'attachment', 'id': identity, 'version': resource['version']},
                'spec': {'edits': [{'target': target['id'], 'old': '12', 'text': '15'}]},
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        output = Presentation(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert [[cell.text for cell in row.cells] for row in output.slides[0].shapes[0].table.rows] == [
            ['Item', 'Count'],
            ['A', '15'],
        ]


def test_excel_revision_changes_one_cell_and_preserves_other_data_and_style(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    workbook = Workbook()
    workbook.active.title = '数据'
    workbook.active.append(['项目', '人数'])
    workbook.active.append(['跑步', 12])
    workbook.active.append(['游泳', 8])
    workbook.active['A2'].font = Font(bold=True)
    data = io.BytesIO()
    workbook.save(data)
    with TestClient(main.app) as client:
        file_id = client.post('/api/work/upload?name=source.xlsx', content=data.getvalue()).json()['task']['id']
        metadata = client.get(f'/api/work/resources/attachment/{file_id}/targets').json()
        response = client.post(
            '/api/work/office/tasks',
            json={
                'operation': 'edit',
                'format': 'xlsx',
                'name': '修订.xlsx',
                'source': {'kind': 'attachment', 'id': file_id, 'version': metadata['version']},
                'spec': {'edits': [{'target': 's:0:B2', 'old': '12', 'value': 15}]},
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        changed = load_workbook(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert changed['数据']['B2'].value == 15
        assert changed['数据']['B3'].value == 8
        assert changed['数据']['A2'].font.bold is True
        assert client.get('/api/work/files/' + file_id).content == data.getvalue()


def test_report_generation_reads_all_source_blocks_in_its_background_task(tmp_path, monkeypatch):
    import json

    from docx import Document

    from backend.app import config, document_models

    captured = []

    async def generate(_value, messages, check=lambda: None, tick=lambda: None):
        check()
        captured.append(str(messages))
        return json.dumps(
            {
                'blocks': [
                    {'type': 'heading', 'text': '完整来源报告', 'level': 0},
                    {'type': 'paragraph', 'text': '已核对开头和终点的材料。'},
                ]
            },
            ensure_ascii=False,
        )

    monkeypatch.setattr(document_models, 'invoke', generate)
    config.save_config(config.ModelConfig(llm_model='offline', llm_api_key='dummy'))
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    source = Document()
    source.add_paragraph('SOURCE_FIRST')
    for index in range(61):
        source.add_paragraph(f'材料第{index}段')
    source.add_paragraph('SOURCE_FINAL')
    data = io.BytesIO()
    source.save(data)
    with TestClient(main.app) as client:
        resource_id = client.post('/api/work/upload?name=source.docx', content=data.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{resource_id}/targets').json()['version']
        created = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'docx',
                'name': '报告.docx',
                'instruction': '根据完整资料制作报告，核对所有段落。',
                'sources': [{'kind': 'attachment', 'id': resource_id, 'version': version}],
            },
        )
        assert created.status_code == 201, created.text
        task = asyncio.run(runtime.work.wait(created.json()['id']))
        assert task['status'] == 'complete', task['error']
        assert any('SOURCE_FIRST' in call and 'SOURCE_FINAL' in call for call in captured)
        report = Document(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert report.paragraphs[0].text == '完整来源报告'


def test_chinese_pdf_embeds_redistributable_font():
    import pypdfium2 as pdfium

    from backend.app.office_tasks import pdf

    data = pdf({'elements': [{'type': 'paragraph', 'text': '中文报告：人数合计二十人，结果可追溯。'}]})
    assert b'/FontFile2' in data
    with pdfium.PdfDocument(data) as document:
        page = document[0]
        text = page.get_textpage()
        try:
            assert '中文报告' in text.get_text_range()
        finally:
            text.close()
            page.close()


def test_word_external_fields_are_blocked_before_the_office_process_starts(tmp_path, monkeypatch):
    from pathlib import Path

    from docx import Document
    from docx.oxml import OxmlElement

    from backend.app import office_runtime

    invoked = []

    def process(*args, **kwargs):
        invoked.append(args)
        raise AssertionError('External field reached Office process')

    monkeypatch.setattr(office_runtime, 'find_office', lambda: Path('unused-office'))
    monkeypatch.setattr(office_runtime.subprocess, 'Popen', process)
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    document = Document()
    field = OxmlElement('w:instrText')
    field.text = ' INCLUDETEXT "file:///blocked-test.invalid/private.txt" '
    document.add_paragraph().add_run()._r.append(field)
    raw = io.BytesIO()
    document.save(raw)
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=unsafe.docx', content=raw.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']
        result = client.post(
            '/api/work/office/tasks',
            json={
                'operation': 'export',
                'format': 'pdf',
                'name': 'unsafe.pdf',
                'source': {'kind': 'attachment', 'id': identity, 'version': version},
            },
        )
        task = asyncio.run(runtime.work.wait(result.json()['id']))
        assert not invoked
        assert task['status'] == 'failed'
        assert '外部字段' in task['error']
        assert not task['context'].get('artifacts')


def test_excel_edit_in_create_shape_without_explicit_source_is_rejected(tmp_path, monkeypatch):
    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        for spec in (
            {'edits': [{'target': 's:0:A1', 'old': 'Name', 'value': ''}]},
            {'sheets': [{'name': 'Sheet1', 'edits': [{'target': 's:0:A1', 'old': 'Name', 'value': ''}]}]},
        ):
            response = client.post(
                '/api/work/office/tasks',
                json={
                    'format': 'xlsx',
                    'name': 'revision.xlsx',
                    'spec': spec,
                },
            )
            assert response.status_code == 422, response.text


def test_excel_revision_validation_rejects_lost_untargeted_cells_and_layout():
    import pytest
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from backend.app.filework import WorkError
    from backend.app.office_excel import edit, verify_edit

    book = Workbook()
    sheet = book.active
    sheet.append(['Vote', 'Name', 'Keep'])
    sheet.append(['Title', 'Candidate', '=1+1'])
    sheet.merge_cells('A4:C4')
    sheet['A4'] = 'Footer'
    sheet['A4'].font = Font(bold=True)
    sheet.column_dimensions['B'].width = 25
    raw = io.BytesIO()
    book.save(raw)
    spec = {'edits': [{'target': 's:0:B2', 'old': 'Candidate', 'value': ''}]}
    correct = edit(raw.getvalue(), spec)
    verify_edit(raw.getvalue(), correct, spec)
    blank = io.BytesIO()
    Workbook().save(blank)
    with pytest.raises(WorkError, match='修订'):
        verify_edit(raw.getvalue(), blank.getvalue(), spec)


def test_saved_artifact_reveal_resolves_id_and_version_before_launch(tmp_path, monkeypatch):
    import subprocess

    from backend.app import work_artifacts

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    launched = []
    monkeypatch.setattr(work_artifacts, 'sys_platform', 'win32', raising=False)
    monkeypatch.setattr(subprocess, 'Popen', lambda args, **kwargs: launched.append((args, kwargs)))
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'xlsx',
                'name': 'safe.xlsx',
                'spec': {'sheets': [{'name': 'Sheet1', 'rows': [['Keep']]}]},
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        version = task['context']['artifact']['sha256']
        base = '/api/work/artifacts/' + task['id'] + '/reveal'
        assert client.post(base, json={'version': '0' * 64}).status_code == 409
        assert launched == []
        assert client.post(base, json={'version': version, 'path': 'C:/other'}).status_code == 422
        assert client.post(base, json={'version': version}).status_code == 200
        assert len(launched) == 1
        args, options = launched[0]
        assert args[1] == '/select,' and args[2].endswith('safe.xlsx')
        assert not options.get('shell')


def test_artifact_thumbnail_is_small_and_rejects_stale_versions(tmp_path, monkeypatch):
    from PIL import Image

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={
                'format': 'pdf',
                'name': 'preview.pdf',
                'spec': {'elements': [{'type': 'paragraph', 'text': 'Preview'}]},
            },
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'complete', task['error']
        version = task['context']['artifact']['sha256']
        base = '/api/work/artifacts/' + task['id'] + '/thumbnail'
        result = client.get(base, params={'version': version})
        assert result.status_code == 200, result.text
        with Image.open(io.BytesIO(result.content)) as image:
            assert max(image.size) <= 240
        assert client.get(base, params={'version': '0' * 64}).status_code == 409


def test_generated_excel_edit_without_source_explains_correction_and_publishes_nothing(tmp_path, monkeypatch):
    import json

    from backend.app import config, document_models

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    config.save_config(config.ModelConfig(llm_model='offline', llm_api_key='dummy'))

    async def generate(*args, **kwargs):
        return json.dumps({'edits': [{'target': 's:0:A1', 'old': 'Candidate', 'value': ''}]})

    monkeypatch.setattr(document_models, 'invoke', generate)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks', json={'format': 'xlsx', 'name': 'revised.xlsx', 'instruction': '只修改姓名'}
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'failed'
        assert 'source' in task['error']
        assert not task['context'].get('artifact')


def test_office_model_timeout_is_not_reported_as_a_disk_or_permission_error(tmp_path, monkeypatch):
    from backend.app import config, document_models

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    config.save_config(config.ModelConfig(llm_model='offline', llm_api_key='dummy'))

    async def timeout(*args, **kwargs):
        raise TimeoutError('PRIVATE_DIAGNOSTIC_MUST_NOT_LEAK')

    monkeypatch.setattr(document_models, 'invoke', timeout)
    with TestClient(main.app) as client:
        response = client.post(
            '/api/work/office/tasks',
            json={'format': 'xlsx', 'name': 'result.xlsx', 'instruction': '根据资料生成工作簿'},
        )
        task = asyncio.run(runtime.work.wait(response.json()['id']))
        assert task['status'] == 'failed'
        assert '超时' in task['error']
        assert '磁盘' not in task['error'] and 'PRIVATE_DIAGNOSTIC' not in task['error']
        assert not task['context'].get('artifact')


def test_pdf_layout_preserving_revision_stops_before_queuing_or_publishing(tmp_path, monkeypatch):
    from reportlab.pdfgen import canvas

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    raw = io.BytesIO()
    pdf = canvas.Canvas(raw)
    pdf.drawString(40, 700, 'Original PDF')
    pdf.save()
    with TestClient(main.app) as client:
        identity = client.post('/api/work/upload?name=original.pdf', content=raw.getvalue()).json()['task']['id']
        version = client.get(f'/api/work/resources/attachment/{identity}/targets').json()['version']
        response = client.post(
            '/api/work/office/tasks',
            json={
                'operation': 'edit',
                'format': 'pdf',
                'name': 'revised.pdf',
                'source': {'kind': 'attachment', 'id': identity, 'version': version},
                'spec': {'edits': [{'target': 'page:1', 'text': 'Changed'}]},
            },
        )
        assert response.status_code == 422, response.text
        assert '排版' in response.json()['detail']
        assert len(client.get('/api/work/tasks').json()) == 1


def test_chat_does_not_convert_a_pdf_to_word_when_the_user_says_not_to(tmp_path, monkeypatch):
    from reportlab.pdfgen import canvas

    from backend.app import feature_config
    from backend.app.tool_registry import ToolRegistry
    from backend.app.work_artifacts import attachment

    store = Store(tmp_path / 'office.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    raw = io.BytesIO()
    pdf = canvas.Canvas(raw)
    pdf.drawString(40, 700, 'Original')
    pdf.save()
    with TestClient(main.app) as client:
        cid = store.create_conversation()['id']
        identity = client.post('/api/work/upload?name=source.pdf', content=raw.getvalue()).json()['task']['id']
        item = attachment(runtime.work, identity)
        message = store.add_message(cid, 'user', '请改错字，保留PDF，不要转Word。', metadata={'attachments': [item]})
        args = {
            'operation': 'edit',
            'format': 'docx',
            'name': 'revised.docx',
            'source': {'kind': 'attachment', 'id': identity, 'version': item['version']},
            'spec': {'blocks': [{'type': 'paragraph', 'text': 'Changed'}]},
        }
        registry = ToolRegistry(feature_config.load(), store, message, [], runtime.work)
        result = asyncio.run(registry.execute('office_files', args))
        assert not result['ok'], result
        assert '排版' in result['message']
        message = store.add_message(cid, 'user', '请将PDF转换为Word。', metadata={'attachments': [item]})
        registry = ToolRegistry(feature_config.load(), store, message, [], runtime.work)
        result = asyncio.run(registry.execute('office_files', args))
        assert result['ok'], result
        task = asyncio.run(runtime.work.wait(result['task']['id']))
        assert task['status'] == 'complete', task['error']
        assert client.get('/api/work/artifacts/' + task['id']).content.startswith(b'PK')
