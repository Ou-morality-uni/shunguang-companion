import asyncio

import pytest
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessageChunk, ToolMessage

from backend.app import config, main, rag, runtime, tools
from backend.app.routes import chat as chat_routes
from backend.app.store import Store


def setup(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'DATA_DIR', tmp_path)
    monkeypatch.setattr(runtime, 'store', Store(tmp_path / 'test.db'))
    monkeypatch.setattr(rag, 'status', lambda: {'ready': False})
    config.save_config(
        config.ModelConfig(llm_base_url='https://example.com/v1', llm_model='test', llm_api_key='test-only')
    )
    runtime.busy.clear()
    runtime.stops.clear()


def test_agent_returns_actual_tool_result_then_persists(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    received = []

    class Model:
        def bind_tools(self, schemas):
            return self

        async def astream(self, history):
            if isinstance(history[-1], ToolMessage):
                received.append(history[-1].content)
                yield AIMessageChunk(content='数据返回：19度。')
            else:
                yield AIMessageChunk(
                    content='', tool_calls=[{'name': 'weather', 'id': 'tool1', 'args': {'city': '上海'}}]
                )

    async def execute(name, args):
        assert name == 'weather' and args == {'city': '上海'}
        return {'ok': True, 'temperature': 19, 'sources': [{'title': 'source', 'url': 'https://example.com/weather'}]}

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    monkeypatch.setattr(tools, 'execute', execute)
    with TestClient(main.app) as client:
        cid = client.post('/api/conversations').json()['id']
        response = client.post('/api/chat', json={'conversation_id': cid, 'content': '上海天气'})
        assert 'event: tool' in response.text and 'event: done' in response.text
        assert '19' in received[0]
        record = client.get(f'/api/conversations/{cid}/messages').json()[-1]
        assert record['tools'][0]['result']['temperature'] == 19
        assert record['status'] == 'complete'


def test_agent_caps_executions_and_rejects_unregistered_tool(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    executed = []

    class Model:
        def bind_tools(self, schemas):
            return self

        async def astream(self, history):
            if len([m for m in history if isinstance(m, ToolMessage)]) >= 10:
                yield AIMessageChunk(content='已到调用上限。')
                return
            yield AIMessageChunk(
                content='',
                tool_calls=[
                    {'name': 'weather' if i < 4 else 'run_shell', 'id': f'{len(history)}-{i}', 'args': {'city': '上海'}}
                    for i in range(5)
                ],
            )

    async def execute(name, args):
        executed.append(name)
        return {'ok': True}

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    monkeypatch.setattr(tools, 'execute', execute)
    with TestClient(main.app) as client:
        cid = client.post('/api/conversations').json()['id']
        response = client.post('/api/chat', json={'conversation_id': cid, 'content': '循环测试'})
        assert response.status_code == 200
        assert executed == ['weather'] * 4
        assert '调用次数已达上限' in response.text


def test_three_sequential_tools_can_finish_before_the_final_reply(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)

    class Model:
        def __init__(self, bound=False):
            self.bound = bound

        def bind_tools(self, schemas):
            return Model(True)

        async def astream(self, history):
            completed = sum(isinstance(item, ToolMessage) for item in history)
            if completed == 3:
                yield AIMessageChunk(content='三个步骤都已完成。')
            elif self.bound:
                yield AIMessageChunk(
                    content='', tool_calls=[{'name': 'weather', 'id': f'chain-{completed}', 'args': {'city': '上海'}}]
                )
            else:
                yield AIMessageChunk(content='工具提前不可用。')

    async def execute(name, args):
        return {'ok': True, 'message': '步骤完成'}

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    monkeypatch.setattr(tools, 'execute', execute)
    with TestClient(main.app) as client:
        cid = client.post('/api/conversations').json()['id']
        client.post('/api/chat', json={'conversation_id': cid, 'content': '依次完成三个步骤'})
        record = client.get(f'/api/conversations/{cid}/messages').json()[-1]
        assert record['content'] == '三个步骤都已完成。'
        assert len(record['tools']) == 3


def test_plain_tool_markup_is_not_streamed_or_treated_as_a_completed_operation(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)

    class Model:
        def bind_tools(self, schemas):
            return self

        async def astream(self, history):
            for part in ['<tool_', 'call><function=office_files>', 'private arguments</tool_call>']:
                yield AIMessageChunk(content=part)

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    with TestClient(main.app) as client:
        cid = client.post('/api/conversations').json()['id']
        response = client.post('/api/chat', json={'conversation_id': cid, 'content': '处理文件'})
        record = client.get(f'/api/conversations/{cid}/messages').json()[-1]
        assert '<tool_' not in response.text and 'private arguments' not in response.text
        assert record['status'] == 'error'
        assert '没有执行' in record['content']
        assert record['tools'] == []


@pytest.mark.parametrize(
    'raw',
    [
        '示例：\n```xml\n<tool_call>example</tool_call>\n```\n以上只是示例。',
        '可以讨论 `<tool_call>` 这个标签。',
        '~~~~xml\n<tool_call>example</tool_call>\n~~~\n<tool_call>still an example</tool_call>\n~~~~',
    ],
)
def test_protocol_examples_remain_readable_in_code_and_ordinary_prose(tmp_path, monkeypatch, raw):
    setup(tmp_path, monkeypatch)

    class Model:
        def bind_tools(self, schemas):
            return self

        async def astream(self, history):
            for char in raw:
                yield AIMessageChunk(content=char)

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    with TestClient(main.app) as client:
        cid = client.post('/api/conversations').json()['id']
        client.post('/api/chat', json={'conversation_id': cid, 'content': '解释协议示例'})
        record = client.get(f'/api/conversations/{cid}/messages').json()[-1]
        assert record['content'] == raw and record['status'] == 'complete'


def test_chat_edits_a_real_ballot_after_guide_and_inspection_without_changing_the_original(tmp_path, monkeypatch):
    import io
    import json

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font

    setup(tmp_path, monkeypatch)
    book = Workbook()
    for column in ('A', 'F'):
        book.active[f'{column}1'] = '保留的票面标题'
        book.active[f'{column}1'].font = Font(bold=True)
        book.active[f'{column}2'] = '姓名'
        book.active[f'{column}3'] = '示例甲'
        book.active[f'{column}4'] = '示例乙'
    buffer = io.BytesIO()
    book.save(buffer)
    original = buffer.getvalue()

    class Model:
        def __init__(self, bound=False):
            self.bound = bound

        def bind_tools(self, schemas):
            return Model(True)

        async def astream(self, history):
            results = [item for item in history if isinstance(item, ToolMessage)]
            if len(results) == 3:
                yield AIMessageChunk(content='修改任务已提交，请查看真实任务结果。')
                return
            if not self.bound:
                yield AIMessageChunk(content='工具不可用，未提交修改。')
                return
            if not results:
                name, args = 'office_guide', {'format': 'xlsx'}
            elif len(results) == 1:
                name, args = 'office_inspect', {'kind': 'attachment', 'id': upload_id}
            else:
                resource = json.loads(results[-1].content)['resource']
                name, args = (
                    'office_files',
                    {
                        'format': 'xlsx',
                        'name': '已清空姓名.xlsx',
                        'operation': 'edit',
                        'source': {'kind': 'attachment', 'id': upload_id, 'version': resource['version']},
                        'spec': {
                            'edits': [
                                {'target': f's:0:{column}{row}', 'old': old, 'value': ''}
                                for column in ('A', 'F')
                                for row, old in ((3, '示例甲'), (4, '示例乙'))
                            ]
                        },
                    },
                )
            yield AIMessageChunk(content='', tool_calls=[{'id': f'office-{len(results)}', 'name': name, 'args': args}])

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    with TestClient(main.app) as client:
        cid = client.post('/api/conversations').json()['id']
        upload_id = client.post('/api/work/upload?name=ballot.xlsx', content=original).json()['task']['id']
        client.post(
            '/api/chat',
            json={'conversation_id': cid, 'content': '清空左右票区的姓名，保留其他内容。', 'attachments': [upload_id]},
        )
        record = client.get(f'/api/conversations/{cid}/messages').json()[-1]
        assert [item['name'] for item in record['tools']] == ['office_guide', 'office_inspect', 'office_files']
        task_id = record['tools'][-1]['result']['task']['id']
        task = asyncio.run(runtime.work.wait(task_id))
        assert task['status'] == 'complete', task['error']
        revised = load_workbook(io.BytesIO(client.get('/api/work/artifacts/' + task_id).content))
        for column in ('A', 'F'):
            assert revised.active[f'{column}3'].value is None and revised.active[f'{column}4'].value is None
            assert revised.active[f'{column}1'].value == '保留的票面标题'
            assert revised.active[f'{column}1'].font.bold is True
            assert revised.active[f'{column}2'].value == '姓名'
        assert client.get('/api/work/files/' + upload_id).content == original


def test_explicit_stop_persists_interruption(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    monkeypatch.setattr(tools, 'schemas', lambda _: [])

    class Model:
        async def astream(self, history):
            yield AIMessageChunk(content='部分')
            await asyncio.sleep(30)

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    runtime.store.initialize()
    cid = runtime.store.create_conversation()['id']

    async def run():
        response = await chat_routes.chat(chat_routes.ChatRequest(conversation_id=cid, content='停止测试'))
        iterator = response.body_iterator
        await anext(iterator)
        await anext(iterator)

        async def drain():
            return [frame async for frame in iterator]

        draining = asyncio.create_task(drain())
        await chat_routes.stop_chat(cid)
        assert runtime.store.messages(cid)[-1]['status'] == 'interrupted'
        remaining = await draining
        assert 'interrupted' in ''.join(remaining)

    asyncio.run(run())
    assert runtime.store.messages(cid)[-1]['content'] == '部分'
    assert runtime.store.messages(cid)[-1]['status'] == 'interrupted'
    assert cid not in runtime.busy


def test_chat_does_not_expose_instant_fact_saving_to_the_main_model(tmp_path, monkeypatch):
    setup(tmp_path, monkeypatch)
    bound = []

    class Model:
        def bind_tools(self, schemas):
            bound.extend(item['name'] for item in schemas)
            return self

        async def astream(self, history):
            yield AIMessageChunk(content='我们接着聊。')

    monkeypatch.setattr(chat_routes, 'build_model', lambda _: Model())
    with TestClient(main.app) as client:
        response = client.post('/api/chat', json={'content': '我喜欢吃苹果。'})
        assert 'event: done' in response.text
        assert bound
        assert 'remember_user_fact' not in bound
        assert 'recall_conversations' in bound
