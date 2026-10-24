import asyncio
import hashlib
import io

from docx import Document
from fastapi.testclient import TestClient

from backend.app import main, runtime
from backend.app.store import Store


def test_annotation_drafts_survive_refresh_and_do_not_modify_the_file(tmp_path, monkeypatch):
    store = Store(tmp_path / 'annotations.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    document = Document()
    document.add_paragraph('要修改的段落')
    data = io.BytesIO()
    document.save(data)
    original = data.getvalue()
    with TestClient(main.app) as client:
        conversation = store.create_conversation()
        upload = client.post('/api/work/upload?name=source.docx', content=original).json()
        file_id = upload['task']['id']
        response = client.post(
            '/api/work/annotations',
            json={
                'conversation_id': conversation['id'],
                'source': {'kind': 'attachment', 'id': file_id, 'version': hashlib.sha256(original).hexdigest()},
                'target_id': 'p:0',
                'quote': '要修改的段落',
                'comment': '请写得更简洁',
            },
        )
        assert response.status_code == 201, response.text
        annotation_id = response.json()['id']
        updated = client.patch('/api/work/annotations/' + annotation_id, json={'comment': '改成一句话'})
        assert updated.status_code == 200
        drafts = client.get('/api/work/annotations', params={'conversation_id': conversation['id']}).json()
        assert drafts[0]['comment'] == '改成一句话'
        assert drafts[0]['status'] == 'draft'
        assert client.get('/api/work/files/' + file_id).content == original
        assert client.delete('/api/work/annotations/' + annotation_id).status_code == 200
        assert client.get('/api/work/annotations', params={'conversation_id': conversation['id']}).json() == []


def test_editing_an_annotation_during_submission_returns_a_conflict_and_keeps_the_new_draft(tmp_path, monkeypatch):
    from backend.app import config
    from backend.app.routes import chat

    store = Store(tmp_path / 'annotations.sqlite3')
    store.initialize()
    monkeypatch.setattr(runtime, 'store', store)
    config.save_config(config.ModelConfig(llm_model='test', llm_api_key='dummy'))
    cid = store.create_conversation()['id']
    original = store.add_message(cid, 'assistant', 'Original answer')
    with TestClient(main.app, raise_server_exceptions=False) as client:
        draft = client.post(
            '/api/work/annotations',
            json={
                'conversation_id': cid,
                'source': {
                    'kind': 'message',
                    'id': original['id'],
                    'version': hashlib.sha256(b'Original answer').hexdigest(),
                },
                'target_id': 'message',
                'quote': 'Original',
                'comment': 'First instruction',
            },
        ).json()

        def connect(_value):
            assert (
                client.patch('/api/work/annotations/' + draft['id'], json={'comment': 'New instruction'}).status_code
                == 200
            )
            return object()

        monkeypatch.setattr(chat, 'build_model', connect)
        response = client.post(
            '/api/chat', json={'conversation_id': cid, 'content': 'Revise the answer.', 'annotations': [draft['id']]}
        )
        assert response.status_code == 409, response.text
        current = client.get('/api/work/annotations', params={'conversation_id': cid}).json()
        assert current[0]['comment'] == 'New instruction'
        assert current[0]['status'] == 'draft'


def test_file_annotation_submission_delivers_a_revision_linked_to_the_comment(tmp_path, monkeypatch):
    import json

    from langchain_core.messages import AIMessageChunk, ToolMessage

    from backend.app import config, document_models
    from backend.app.routes import chat

    config.save_config(config.ModelConfig(llm_model='test', llm_api_key='dummy'))
    store = Store(tmp_path / 'annotations.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    document = Document()
    document.add_paragraph('Original report')
    data = io.BytesIO()
    document.save(data)
    reference = {}
    generator_messages = []

    async def generate(_value, history, check=lambda: None, tick=lambda: None):
        check()
        generator_messages.append(str(history))
        return json.dumps({'edits': [{'target': 'p:0', 'old': 'Original', 'text': 'Revised'}]})

    monkeypatch.setattr(document_models, 'invoke', generate)

    class Model:
        def bind_tools(self, *args, **kwargs):
            return self

        async def astream(self, history):
            if isinstance(history[-1], ToolMessage):
                yield AIMessageChunk(content='修订任务已经安排，完成后可查看新文件。')
            else:
                yield AIMessageChunk(
                    content='',
                    tool_calls=[
                        {
                            'id': 'revision',
                            'name': 'office_files',
                            'args': {
                                'format': 'docx',
                                'name': 'revised.docx',
                                'operation': 'edit',
                                'source': reference,
                                'instruction': '按本次提交的批注修订文件。',
                            },
                        }
                    ],
                )

    monkeypatch.setattr(chat, 'build_model', lambda _value: Model())
    with TestClient(main.app) as client:
        cid = store.create_conversation()['id']
        identity = client.post('/api/work/upload?name=source.docx', content=data.getvalue()).json()['task']['id']
        reference.update(kind='attachment', id=identity, version=hashlib.sha256(data.getvalue()).hexdigest())
        draft = client.post(
            '/api/work/annotations',
            json={
                'conversation_id': cid,
                'source': reference,
                'target_id': 'p:0',
                'quote': 'Original',
                'comment': '把 Original 改成 Revised',
            },
        ).json()
        sent = client.post(
            '/api/chat', json={'conversation_id': cid, 'content': '按批注修改', 'annotations': [draft['id']]}
        )
        assert 'event: done' in sent.text
        jobs = client.get('/api/work/tasks', params={'conversation_id': cid}).json()
        task = asyncio.run(runtime.work.wait(jobs[0]['id']))
        assert task['status'] == 'complete', task['error']
        assert any('把 Original 改成 Revised' in messages for messages in generator_messages)
        result = task['steps'][-1]['result']
        assert result['annotation_results'] == [{'id': draft['id'], 'location': '正文第 1 段', 'status': 'applied'}]
        changed = Document(io.BytesIO(client.get('/api/work/artifacts/' + task['id']).content))
        assert changed.paragraphs[0].text == 'Revised report'
        assert client.get('/api/work/files/' + identity).content == data.getvalue()


def test_submitting_reply_annotations_adds_complete_revision_and_keeps_original(tmp_path, monkeypatch):
    from langchain_core.messages import AIMessageChunk

    from backend.app import config
    from backend.app.routes import chat

    captured = []

    class Model:
        def bind_tools(self, *args, **kwargs):
            return self

        async def astream(self, messages, **kwargs):
            captured.extend(messages)
            yield AIMessageChunk(content='这是按批注生成的完整修订回复。')

    monkeypatch.setattr(chat, 'build_model', lambda _: Model())
    config.save_config(config.ModelConfig(llm_model='offline', llm_api_key='dummy'))
    store = Store(tmp_path / 'annotations.sqlite3')
    monkeypatch.setattr(runtime, 'store', store)
    with TestClient(main.app) as client:
        conversation = store.create_conversation()
        original = store.add_message(conversation['id'], 'assistant', '原回复的具体内容。')
        preview = client.get(f'/api/work/resources/message/{original["id"]}/targets')
        assert preview.status_code == 200
        assert preview.json()['targets'][0]['id'] == 'message'
        annotation = client.post(
            '/api/work/annotations',
            json={
                'conversation_id': conversation['id'],
                'source': {
                    'kind': 'message',
                    'id': original['id'],
                    'version': hashlib.sha256(original['content'].encode()).hexdigest(),
                },
                'target_id': 'message',
                'quote': '具体内容',
                'comment': '解释得更清楚一些',
            },
        ).json()
        response = client.post(
            '/api/chat',
            json={
                'conversation_id': conversation['id'],
                'content': '请按批注修订。',
                'annotations': [annotation['id']],
            },
        )
        assert response.status_code == 200, response.text
        messages = client.get(f'/api/conversations/{conversation["id"]}/messages').json()
        assert messages[0]['content'] == '原回复的具体内容。'
        assert messages[-1]['content'] == '这是按批注生成的完整修订回复。'
        assert messages[-1]['annotations'][0]['id'] == annotation['id']
        assert any('解释得更清楚一些' in str(message.content) for message in captured)
        assert client.get('/api/work/annotations', params={'conversation_id': conversation['id']}).json() == []
        assert (
            client.patch(
                '/api/work/annotations/' + annotation['id'], json={'comment': '不能重写已发送批注'}
            ).status_code
            == 409
        )
