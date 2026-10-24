"""Versioned, read-only object identities for previews and targeted edits."""

import hashlib
import io
import zipfile
from collections import OrderedDict
from pathlib import Path
from threading import Lock

from . import document_readers
from .filework import WorkError
from .work_artifacts import artifact_bytes, attachment_bytes

PREVIEWS = OrderedDict()
PREVIEW_LOCK = Lock()


def belongs_to_chat(work, conversation_id, kind, identity):
    if not conversation_id:
        return False
    record = work.task(identity.split('~', 1)[0])
    if kind == 'artifact':
        return record.get('conversation_id') == conversation_id
    return (
        record.get('context', {}).get('imported_conversation_id') == conversation_id
        and bool(record.get('context', {}).get('imported_from'))
    ) or any(
        item.get('id') == identity
        for message in work.store.messages(conversation_id)
        for item in message.get('attachments', [])
    )


def import_file(work, workspace_id, path, conversation_id):
    with work.gate:
        root = work.root(workspace_id)
        source = work.path(root, path)
        if source.suffix.lower() not in (
            '.docx',
            '.xlsx',
            '.pptx',
            '.pdf',
            '.png',
            '.jpg',
            '.jpeg',
            '.webp',
            '.bmp',
            '.gif',
            '.txt',
            '.md',
        ):
            raise WorkError('请使用常规办公文档、PDF、图片或文本；旧 Office 格式先另存为常规格式。', 422)
        stamp = work.fingerprint(source)
        if stamp is None:
            raise WorkError('原件不存在。', 404)
        raw = source.read_bytes()
        work.check_state(root, {path: stamp})
        uploaded = work.upload(source.name, raw)
        original = {'workspace_id': workspace_id, 'path': path, 'version': stamp['sha256']}
        record = work.task(uploaded['task']['id'])
        record['context']['imported_from'] = original
        record['context']['imported_conversation_id'] = conversation_id
        work.store.save_file_task(record)
        return {
            'ok': True,
            'source': {'kind': 'attachment', 'id': record['id'], 'version': stamp['sha256']},
            'original': original,
            'name': source.name,
            'message': '已从授权目录读取并保存本次处理副本；原件未修改。仅在用户明确要求覆盖时使用 original，并等待覆盖确认。',
        }


def preview_pdf(work, kind, resource_id, version=None):
    item, data = resource_bytes(work, kind, resource_id, version)
    suffix = Path(item['name']).suffix.lower()
    if suffix == '.pdf':
        return data
    if suffix not in ('.docx', '.xlsx', '.pptx'):
        raise WorkError('此资源不支持页面排版预览。', 422)
    key = item['version'] + suffix
    with PREVIEW_LOCK:
        if key in PREVIEWS:
            PREVIEWS.move_to_end(key)
            return PREVIEWS[key]
        from .office_runtime import convert

        converted = convert(data, suffix)
        resource_bytes(work, kind, resource_id, item['version'])
        PREVIEWS[key] = converted
        while len(PREVIEWS) > 10 or sum(len(value) for value in PREVIEWS.values()) > 64 * 1024 * 1024:
            PREVIEWS.popitem(last=False)
        return converted


def page_info(work, kind, resource_id, version=None):
    import pypdfium2 as pdfium

    data = preview_pdf(work, kind, resource_id, version)
    with document_readers.PDF_LOCK, pdfium.PdfDocument(data) as document:
        return {'pages': len(document)}


def page_image(work, kind, resource_id, page=1, version=None):
    import pypdfium2 as pdfium

    item, data = resource_bytes(work, kind, resource_id, version)
    if Path(item['name']).suffix.lower() not in ('.docx', '.xlsx', '.pptx', '.pdf'):
        try:
            return document_readers.image_bytes(data)
        except document_readers.DocumentError as error:
            raise WorkError(str(error), 422)
    data = preview_pdf(work, kind, resource_id, item['version'])
    with document_readers.PDF_LOCK, pdfium.PdfDocument(data) as document:
        if not 1 <= page <= len(document):
            raise WorkError('预览页码超出范围。', 422)
        source = document[page - 1]
        bitmap = None
        try:
            width, height = source.get_size()
            if width <= 0 or height <= 0:
                raise WorkError('页面尺寸无效。', 422)
            bitmap = source.render(scale=min(1.8, 1600 / max(width, height)))
            output = io.BytesIO()
            bitmap.to_pil().convert('RGB').save(output, format='JPEG')
            return output.getvalue()
        finally:
            if bitmap:
                bitmap.close()
            source.close()


def resource_bytes(work, kind, resource_id, version=None):
    if kind not in ('attachment', 'artifact'):
        raise WorkError('资源类型无效。', 422)
    item, data = (attachment_bytes if kind == 'attachment' else artifact_bytes)(work, resource_id)
    actual = hashlib.sha256(data).hexdigest()
    if version is not None and actual != version:
        raise WorkError('原文件版本已变化，请重新预览并定位后再修改。')
    return {**item, 'version': actual}, data


def targets(work, kind, resource_id):
    if kind == 'message':
        with work.store.connect() as db:
            row = db.execute(
                "SELECT * FROM messages WHERE id=? AND role='assistant' AND status='complete'", (resource_id,)
            ).fetchone()
        if not row:
            raise WorkError('回复不存在或尚未完成。', 404)
        return {
            'id': resource_id,
            'kind': kind,
            'name': '小光的回复',
            'version': hashlib.sha256(row['content'].encode()).hexdigest(),
            'targets': [{'id': 'message', 'kind': 'paragraph', 'text': row['content'], 'location': '聊天回复'}],
            'truncated': False,
        }
    item, data = resource_bytes(work, kind, resource_id)
    suffix = Path(item['name']).suffix.lower()
    values = []
    remaining, truncated = 120000, False

    def add(identity, text, location, **properties):
        nonlocal remaining, truncated
        if len(values) >= 2000 or remaining <= 0:
            truncated = True
            return
        raw = str(text)
        clipped = raw[: min(remaining, 20000)]
        truncated = truncated or len(clipped) < len(raw)
        remaining -= len(clipped)
        values.append({'id': identity, 'text': clipped, 'location': location, **properties})

    if suffix in ('.docx', '.xlsx', '.pptx'):
        try:
            document_readers.check_archive(data)
        except (ValueError, zipfile.BadZipFile, document_readers.DocumentError) as error:
            raise WorkError('文件结构不受支持，无法提供可编辑目标。', 422) from error
    if suffix == '.docx':
        from docx import Document

        document = Document(io.BytesIO(data))
        for index, paragraph in enumerate(document.paragraphs):
            add(f'p:{index}', paragraph.text, f'正文第 {index + 1} 段', kind='paragraph')
        for section_index, section in enumerate(document.sections):
            for prefix, label, container in [('h', '页眉', section.header), ('f', '页脚', section.footer)]:
                if not container.is_linked_to_previous:
                    for index, paragraph in enumerate(container.paragraphs):
                        add(
                            f'{prefix}:{section_index}:{index}',
                            paragraph.text,
                            f'第 {section_index + 1} 节{label}第 {index + 1} 段',
                            kind='paragraph',
                        )
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                for column, cell in enumerate(row.cells):
                    add(
                        f't:{table_index}:{row_index}:{column}',
                        cell.text,
                        f'表格 {table_index + 1} 第 {row_index + 1} 行第 {column + 1} 列',
                        kind='table_cell',
                    )
    elif suffix == '.xlsx':
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet_index, sheet in enumerate(workbook):
                if sheet.max_row > 100000 or sheet.max_column > 1000 or sheet.max_row * sheet.max_column > 200000:
                    raise WorkError('工作表范围过大，请先缩小文件或任务范围。', 422)
                for row in sheet.iter_rows():
                    if len(values) >= 2000:
                        break
                    for cell in row:
                        if cell.value is not None:
                            add(
                                f's:{sheet_index}:{cell.coordinate}',
                                cell.value,
                                f'{sheet.title} · {cell.coordinate}',
                                kind='cell',
                            )
        finally:
            workbook.close()
    elif suffix == '.pptx':
        from pptx import Presentation

        deck = Presentation(io.BytesIO(data))
        for slide_index, slide in enumerate(deck.slides):
            for shape in slide.shapes:
                properties = {
                    'object_type': 'text'
                    if shape.has_text_frame
                    else 'table'
                    if shape.has_table
                    else 'chart'
                    if shape.has_chart
                    else 'image'
                }
                if shape.has_chart:
                    chart = shape.chart
                    try:
                        categories = list(chart.plots[0].categories)
                        properties['chart'] = {
                            'categories': [category.label for category in categories[:100]],
                            'series': [
                                {'name': series.name, 'values': list(series.values)[:100]}
                                for series in list(chart.series)[:8]
                            ],
                            'truncated': len(categories) > 100 or len(chart.series) > 8,
                        }
                    except (AttributeError, ValueError, TypeError, NotImplementedError):
                        properties['chart'] = {'unsupported': '此复杂图表暂不能提供常规类别数据，请在原文件中核对。'}
                text = (
                    shape.text
                    if shape.has_text_frame
                    else (
                        '\n'.join(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                        if shape.has_table
                        else shape.name
                    )
                )
                add(
                    f'slide:{slide_index}:{shape.shape_id}',
                    text,
                    f'第 {slide_index + 1} 页 · {shape.name}',
                    kind='shape',
                    page=slide_index + 1,
                    bounds=[shape.left, shape.top, shape.width, shape.height],
                    canvas=[deck.slide_width, deck.slide_height],
                    **properties,
                )
                if shape.has_table:
                    for row_index, row in enumerate(shape.table.rows):
                        for column, cell in enumerate(row.cells):
                            add(
                                f'slide:{slide_index}:{shape.shape_id}:cell:{row_index}:{column}',
                                cell.text,
                                f'第 {slide_index + 1} 页 · {shape.name} · 第 {row_index + 1} 行第 {column + 1} 列',
                                kind='table_cell',
                                page=slide_index + 1,
                            )
    elif suffix == '.pdf':
        import pypdfium2 as pdfium

        with document_readers.PDF_LOCK, pdfium.PdfDocument(data) as document:
            truncated = truncated or len(document) > 200
            for index in range(min(len(document), 200)):
                page = document[index]
                textpage = page.get_textpage()
                try:
                    add(
                        f'page:{index + 1}',
                        textpage.get_text_bounded(),
                        f'第 {index + 1} 页',
                        kind='page',
                        page=index + 1,
                        canvas=list(page.get_size()),
                    )
                finally:
                    textpage.close()
                    page.close()
    elif suffix in ('.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tiff', '.gif'):
        from PIL import Image, ImageOps

        with Image.open(io.BytesIO(data)) as image:
            image = ImageOps.exif_transpose(image)
            add('image', item['name'], '整张图片', kind='image', canvas=[image.width, image.height])
    elif suffix in ('.md', '.txt'):
        text = data.decode('utf-8')
        values.append({'id': 'document', 'text': text[:60000], 'location': '文档正文', 'kind': 'paragraph'})
    else:
        raise WorkError('此格式暂不支持定位批注，请使用常规 Office、PDF、图片或 Markdown。', 422)
    return {
        'id': resource_id,
        'kind': kind,
        'name': item['name'],
        'version': item['version'],
        'targets': values,
        'truncated': truncated or len(values) >= 2000,
        'validation': item.get('validation', {}),
    }
