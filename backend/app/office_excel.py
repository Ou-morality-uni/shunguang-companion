"""Targeted workbook changes; formulas remain formulas, never hidden as values."""

import io
import re
import warnings
import zipfile

from .filework import WorkError


def calculated(data, check):
    from openpyxl import load_workbook

    from .office_runtime import convert, safe_package

    safe_package(data)
    workbook = load_workbook(io.BytesIO(data), keep_links=False)
    if sum(sheet.max_row * sheet.max_column for sheet in workbook) > 100000:
        raise WorkError('工作簿范围过大，请缩小到十万个单元格内再计算。', 422)
    formulas = {}
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                check()
                if cell.data_type == 'f':
                    validate_formula(cell.value)
                    formulas[(sheet.title, cell.coordinate)] = cell.value
    if not formulas:
        return data, {'calculation': 'not_needed'}
    result = convert(data, '.xlsx', 'xlsx', check)
    computed = load_workbook(io.BytesIO(result), data_only=True, keep_links=False)
    retained = load_workbook(io.BytesIO(result), data_only=False, keep_links=False)
    if workbook.sheetnames != retained.sheetnames:
        raise WorkError('重算改变了工作表结构，未发布可能丢失内容的副本。', 422)
    for (sheet, coordinate), formula in formulas.items():
        check()
        cached = computed[sheet][coordinate]
        if retained[sheet][coordinate].value != formula or cached.value is None or cached.data_type == 'e':
            raise WorkError(f'{sheet}!{coordinate} 的公式未保留、未获得计算结果或存在错误；请检查公式兼容性。', 422)

    def objects(raw):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            return sum(name.startswith('xl/charts/chart') and name.endswith('.xml') for name in archive.namelist())

    if objects(data) != objects(result):
        raise WorkError('重算可能丢失原生图表，未发布降级副本。', 422)
    return result, {'calculation': 'passed', 'formula_cells': len(formulas), 'engine': 'LibreOffice'}


def validate_formula(value):
    if not isinstance(value, str) or not value.startswith('=') or len(value) > 8000:
        raise WorkError('公式格式无效。', 422)
    compact = re.sub(r'\s+', '', value.upper())
    if any(
        token in compact
        for token in ('[', '|', 'WEBSERVICE(', 'HYPERLINK(', 'RTD(', 'DDE(', 'IMAGE(', 'CALL(', 'REGISTER.ID(', 'EXEC(')
    ):
        raise WorkError('不执行外部链接、命令或数据连接公式。', 422)


def edit(data, spec):
    from openpyxl import load_workbook

    edits = spec.get('edits', [])
    if not isinstance(edits, list) or not 1 <= len(edits) <= 1000:
        raise WorkError('一次 Excel 修订需要 1～1000 条单元格修改。', 422)
    with warnings.catch_warnings(record=True) as notices:
        warnings.simplefilter('always')
        workbook = load_workbook(io.BytesIO(data), keep_links=True)
    if any(
        'not supported' in str(item.message).lower() or 'will be removed' in str(item.message).lower()
        for item in notices
    ):
        raise WorkError('此工作簿包含暂不能可靠保留的复杂对象，请另存常规副本后再编辑。', 422)
    touched = set()
    for change in edits:
        target = str(change.get('target', ''))
        match = re.fullmatch(r's:([0-9]{1,3}):([A-Z]{1,3}[1-9][0-9]{0,5})', target)
        if not match or target in touched:
            raise WorkError('单元格位置无效或被重复修改。', 422)
        touched.add(target)
        sheet_index = int(match[1])
        if sheet_index >= len(workbook.worksheets):
            raise WorkError('工作表不存在。', 404)
        cell = workbook.worksheets[sheet_index][match[2]]
        if cell.row > 20000 or cell.column > 100:
            raise WorkError('单元格超出本次编辑范围，请缩小文件。', 422)
        if str(cell.value) != str(change.get('old')):
            raise WorkError('原单元格内容已变化或不匹配，请重新定位。')
        if 'formula' in change:
            validate_formula(change['formula'])
            cell.value = change['formula']
        elif 'value' in change:
            value = change['value']
            if value is not None and not isinstance(value, (str, int, float, bool)):
                raise WorkError('单元格内容须为文字、数字、布尔值或空值。', 422)
            cell.value = value
            if isinstance(value, str) and value.startswith('='):
                cell.data_type = 's'
        else:
            raise WorkError('修改必须明确提供 value 或 formula。', 422)
    workbook.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    workbook.save(output)
    result = output.getvalue()
    verify_edit(data, result, spec)
    return result


def verify_edit(original, revised, spec, layout=True):
    """Do not publish a readable workbook that silently lost unrelated content."""
    from copy import copy

    from openpyxl import load_workbook

    before = load_workbook(io.BytesIO(original), keep_links=True)
    after = load_workbook(io.BytesIO(revised), keep_links=True)
    expected = {}
    for change in spec['edits']:
        _, index, coordinate = change['target'].split(':')
        value = change.get('formula', change.get('value'))
        expected[(int(index), coordinate)] = None if value == '' else value
    failure = 'Excel 修订结果与指定修改不一致，或改变了非目标内容/版式；未发布此副本。'
    if before.sheetnames != after.sheetnames:
        raise WorkError(failure, 422)
    for index, (left, right) in enumerate(zip(before, after)):
        rows, columns = max(left.max_row, right.max_row), max(left.max_column, right.max_column)
        if rows * columns > 100000:
            raise WorkError('Excel 修订验证范围超过十万个单元格，请缩小文件。', 422)
        for row in left.iter_rows(max_row=rows, max_col=columns):
            for cell in row:
                changed = (index, cell.coordinate) in expected
                value = expected.get((index, cell.coordinate), cell.value)
                actual = right[cell.coordinate]
                if actual.value != value or (not changed and actual.data_type != cell.data_type):
                    raise WorkError(failure, 422)
                if layout and any(
                    copy(getattr(cell, field)) != copy(getattr(actual, field))
                    for field in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection')
                ):
                    raise WorkError(failure, 422)
        if str(left.merged_cells) != str(right.merged_cells):
            raise WorkError(failure, 422)
        if layout:
            for name, dimension in left.column_dimensions.items():
                target = right.column_dimensions[name]
                if (dimension.width, dimension.hidden) != (target.width, target.hidden):
                    raise WorkError(failure, 422)
            for name, dimension in left.row_dimensions.items():
                target = right.row_dimensions[name]
                if (dimension.height, dimension.hidden) != (target.height, target.hidden):
                    raise WorkError(failure, 422)
