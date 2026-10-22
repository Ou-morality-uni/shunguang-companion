"""Bounded, read-only extraction. Never follows Office relationships to the network."""

import io
import zipfile
from pathlib import Path
from threading import Lock

MAX_CONTENT = 60000
MAX_VISUALS = 8
PDF_LOCK = Lock()  # PDFium may not be called concurrently, even for different documents.


class DocumentError(Exception):
    pass


def image_bytes(data):
    from PIL import Image, ImageOps

    with Image.open(io.BytesIO(data)) as image:
        if image.width * image.height > 16000000:
            raise DocumentError('图片像素过大，最多支持 1600 万像素。')
        image = ImageOps.exif_transpose(image).convert('RGB')
        image.thumbnail((1600, 1600))
        output = io.BytesIO()
        image.save(output, format='JPEG', quality=85)
        return output.getvalue()


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 4000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
            raise DocumentError('办公文档解压后过大，无法在本次任务中安全读取。')
        for entry in entries:
            if entry.flag_bits & 1:
                raise DocumentError('文档受到加密保护，请先提供可读取的副本。')
            if entry.filename.endswith(('.xml', '.rels')):
                raw = archive.read(entry)
                if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                    raise DocumentError('文档包含不支持的 XML 实体定义。')


class Extraction:
    def __init__(self, name, start, limit, check, force_visual=False):
        self.result = {'file': name, 'blocks': [], 'warnings': [], 'truncated': False, 'next_start': None}
        self.visuals = []
        self.start, self.limit, self.check = start, limit, check
        self.size = 0
        self.force_visual = force_visual

    def window(self, total, unit):
        self.result.update(total_units=total, unit=unit, start=self.start)
        if self.start > max(total, 1):
            raise DocumentError('起始位置超出文档范围。')
        end = min(total, self.start - 1 + self.limit)
        if self.start > 1:
            self.result['truncated'] = True
            self.warn(f'本次从第 {self.start} {unit}开始，之前的范围未读取。')
        if end < total:
            self.result.update(truncated=True, next_start=end + 1)
            self.warn('只读取指定范围，不能把范围外的内容当作已经读取。')
        return range(self.start - 1, end)

    def warn(self, text):
        if text not in self.result['warnings']:
            self.result['warnings'].append(text)

    def add(self, location, text, kind='text'):
        self.check()
        if not text:
            return
        available = MAX_CONTENT - self.size
        if len(text) > available:
            self.result['truncated'] = True
            self.result['next_start'] = None
            self.warn('提取文字已达 60000 字符上限，部分内容未返回；请依据最后返回的位置缩小范围重新读取。')
        text = text[: max(0, available)]
        self.size += len(text)
        if text:
            self.result['blocks'].append({'location': location, 'kind': kind, 'text': text})

    def image(self, location, data, strict=False):
        self.check()
        if len(self.visuals) >= MAX_VISUALS:
            self.result['truncated'] = True
            self.warn('本次最多识别 8 幅图像，其余图像尚未识别。')
            return
        try:
            encoded = image_bytes(data)
        except (OSError, DocumentError):
            if strict:
                raise DocumentError('图片格式、尺寸或内容不受支持。')
            self.result['truncated'] = True
            self.warn(f'{location}的图像格式或尺寸不受支持，未识别该图像。')
            return
        self.visuals.append({'location': location, 'data': encoded})


def pdf(data, extraction):
    import pypdfium2 as pdfium

    with PDF_LOCK, pdfium.PdfDocument(data) as document:
        for index in extraction.window(len(document), '页'):
            extraction.check()
            page = document[index]
            textpage = page.get_textpage()
            try:
                location = f'第 {index + 1} 页'
                if textpage.count_chars() > 100000:
                    raise DocumentError('单页文字过多，请使用较小的文档副本。')
                text = textpage.get_text_bounded()
                extraction.add(location, text)
                has_image = next(page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_IMAGE]), None) is not None
                if extraction.force_visual or not text.strip() or has_image:
                    width, height = page.get_size()
                    if width <= 0 or height <= 0:
                        raise DocumentError('PDF 页面尺寸无效。')
                    bitmap = page.render(scale=min(1.8, 1600 / max(width, height)))
                    try:
                        output = io.BytesIO()
                        bitmap.to_pil().save(output, format='PNG')
                        extraction.image(location, output.getvalue())
                    finally:
                        bitmap.close()
            finally:
                textpage.close()
                page.close()
        extraction.warn(
            'PDF 的文字顺序可能与视觉排版不同；视觉识别结果也应结合原页核对。'
            if extraction.force_visual
            else 'PDF 的文字顺序可能与视觉排版不同；仅有文字的矢量图表未作完整视觉解读，可选择页面视觉理解。'
        )


def word(data, extraction):
    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(data))
    blocks = list(document.iter_inner_content())
    for index in extraction.window(len(blocks), '段落／表格块'):
        block = blocks[index]
        location = f'正文第 {index + 1} 块'
        if isinstance(block, Table):
            extraction.add(
                location, '\n'.join(' | '.join(cell.text for cell in row.cells) for row in block.rows), 'table'
            )
        else:
            extraction.add(location, block.text)
        for number, blip in enumerate(block._element.xpath('.//a:blip'), 1):
            relationship = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            if relationship and relationship in document.part.related_parts:
                extraction.image(f'{location}，图片 {number}', document.part.related_parts[relationship].blob)
    for section_number, section in enumerate(document.sections, 1):
        for label, container in (('页眉', section.header), ('页脚', section.footer)):
            extraction.add(f'第 {section_number} 节{label}', '\n'.join(item.text for item in container.paragraphs))
    extraction.warn('Word 来源按段落／表格定位；未分页渲染，页码、修订、文本框和复杂嵌入对象不保证完整提取。')


def excel(data, extraction):
    from openpyxl import load_workbook

    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        for index in extraction.window(len(formulas.worksheets), '工作表'):
            sheet, cached = formulas.worksheets[index], values.worksheets[index]
            count = 0
            for row_number, (row, cached_row) in enumerate(
                zip(sheet.iter_rows(max_row=500, max_col=100), cached.iter_rows(max_row=500, max_col=100)), 1
            ):
                extraction.check()
                cells = []
                for cell, cached_cell in zip(row, cached_row):
                    if cell.value is None:
                        continue
                    count += 1
                    text = str(cell.value)
                    if cell.data_type == 'f':
                        text += f' [已保存计算值：{cached_cell.value if cached_cell.value is not None else "无缓存，未重新计算"}]'
                    cells.append(f'{cell.coordinate}={text}')
                if cells:
                    extraction.add(f'工作表「{sheet.title}」第 {row_number} 行', ' | '.join(cells), 'cells')
                if count >= 5000 or extraction.size >= MAX_CONTENT:
                    extraction.result['truncated'] = True
                    extraction.warn('工作表读取已达单次 5000 个非空单元格或文字上限。')
                    break
            if (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 100:
                extraction.result['truncated'] = True
                extraction.warn('每个工作表最多读取前 500 行、100 列，范围外内容未处理。')
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if extraction.start == 1:
                for item in archive.namelist():
                    if item.startswith('xl/media/') and not item.endswith('/'):
                        extraction.image(f'工作簿内嵌图片 {Path(item).name}（未定位到单元格）', archive.read(item))
        extraction.warn('公式只读取表达式与已保存缓存值，不执行公式、不刷新外部数据，也不保证图表视觉布局已理解。')
    finally:
        values.close()
        formulas.close()


def powerpoint(data, extraction):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(data))

    def shapes(items, location):
        for number, shape in enumerate(items, 1):
            extraction.check()
            where = f'{location}，对象 {number}'
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                shapes(shape.shapes, where)
            if shape.has_text_frame:
                extraction.add(where, shape.text)
            if shape.has_table:
                extraction.add(
                    where, '\n'.join(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows), 'table'
                )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                extraction.image(where, shape.image.blob)
            if shape.has_chart:
                try:
                    categories = [category.label for category in shape.chart.plots[0].categories]
                    text = (
                        '类别：'
                        + str(categories)
                        + '\n'
                        + '\n'.join(f'{series.name}: {list(series.values)}' for series in shape.chart.series)
                    )
                    extraction.add(where, text, 'chart_values')
                except (ValueError, AttributeError):
                    extraction.warn(f'{where}的图表数据未能完整提取。')

    for index in extraction.window(len(presentation.slides), '幻灯片'):
        slide = presentation.slides[index]
        location = f'第 {index + 1} 张幻灯片'
        shapes(slide.shapes, location)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame:
                extraction.add(f'{location}，备注', frame.text)
    extraction.warn('提取文字、表格、图表数值和内嵌图片；未渲染整页布局、动画或复杂嵌入对象。')


def extract(data, name, start=1, limit=20, check=lambda: None, force_visual=False):
    extraction = Extraction(name, start, limit, check, force_visual)
    suffix = Path(name).suffix.lower()
    try:
        if suffix == '.pdf':
            pdf(data, extraction)
        elif suffix in ('.docx', '.xlsx', '.pptx'):
            check_archive(data)
            {'.docx': word, '.xlsx': excel, '.pptx': powerpoint}[suffix](data, extraction)
        elif suffix in ('.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff', '.gif'):
            extraction.window(1, '图片')
            extraction.image('图片第 1 帧', data, strict=True)
            extraction.warn('多帧图片只处理第一帧。识字结果不等于完整理解所有图形。')
        elif suffix in ('.doc', '.xls', '.ppt', '.docm', '.xlsm', '.pptm'):
            raise DocumentError('请先另存为常规 docx／xlsx／pptx；本版不执行宏或自动转换旧格式。')
        else:
            text = data.decode('utf-8-sig')
            if '\x00' in text:
                raise DocumentError('文件不是可直接读取的 UTF-8 文本，也不属于支持的文档格式。')
            lines = text.splitlines()
            for index in extraction.window(len(lines), '代码／文本行'):
                extraction.add(
                    f'第 {index + 1} 行',
                    lines[index],
                    'code' if suffix not in ('.txt', '.md', '.csv', '.tsv') else 'text',
                )
    except DocumentError:
        raise
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, UnicodeError) as error:
        raise DocumentError('文档无法读取，可能已加密、损坏或格式不受支持。') from error
    return extraction.result, extraction.visuals
